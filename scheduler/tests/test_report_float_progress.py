from datetime import date, timedelta

from django.test import TestCase
from django.utils import timezone

from scheduler import report_export, report_service
from scheduler.models import Project, ScheduleUpload


def _act(code, **extra):
    d = {
        'id': code, 'code': code, 'name': f'Activity {code}', 'wbs': 'Area A', 'area': 'Area A', 'discipline': 'Electrical',
        'bStart': '2026-01-01', 'bFinish': '2026-01-20', 'dur': 20.0, 'remainDur': 20.0, 'pctComplete': 0.0,
        'totalFloat': 0.0, 'isCritical': False, 'onLongestPath': False, 'isMilestone': False,
        'predecessors': [], 'successors': [],
    }
    d.update(extra)
    return d


class FloatIntelligenceSectionTests(TestCase):
    def setUp(self):
        self.project = Project.objects.create(name='Float Intelligence Report Test')
        now = timezone.now()
        self.baseline = ScheduleUpload.objects.create(
            project=self.project, original_filename='bl.xer', sanitized_filename='bl.xer', file_type='XER',
            schedule_classification='APPROVED_BASELINE', data_date=date(2026, 1, 1), version_label='Baseline',
            activities_json=[_act('A1', totalFloat=20.0)], upload_timestamp=now - timedelta(days=30),
        )
        self.curr = ScheduleUpload.objects.create(
            project=self.project, original_filename='c.xer', sanitized_filename='c.xer', file_type='XER',
            schedule_classification='CURRENT_UPDATE', data_date=date(2026, 8, 15), version_label='Current',
            activities_json=[_act('A1', totalFloat=-5.0, isCritical=True)], upload_timestamp=now,
        )

    def test_available_with_real_deterioration(self):
        section = report_service.build_float_intelligence_section(self.project, self.curr)
        self.assertTrue(section['available'])
        self.assertEqual(section['summary']['negativeFloatCount'], 1)
        self.assertEqual(len(section['topDeterioration']), 1)

    def test_unavailable_without_current_version(self):
        section = report_service.build_float_intelligence_section(self.project, None)
        self.assertFalse(section['available'])

    def test_never_truncates_deterioration_list(self):
        # Isolated project — avoids ambiguity with setUp()'s own
        # baseline/current versions when resolving "the" baseline version.
        project = Project.objects.create(name='Float Intelligence No-Truncation Test')
        acts = [_act(f'A{i}', totalFloat=float(-i)) for i in range(1, 40)]
        baseline_acts = [_act(f'A{i}', totalFloat=20.0) for i in range(1, 40)]
        ScheduleUpload.objects.create(
            project=project, original_filename='bl2.xer', sanitized_filename='bl2.xer', file_type='XER',
            schedule_classification='APPROVED_BASELINE', data_date=date(2026, 1, 1),
            activities_json=baseline_acts, upload_timestamp=timezone.now() - timedelta(days=60),
        )
        curr = ScheduleUpload.objects.create(
            project=project, original_filename='c2.xer', sanitized_filename='c2.xer', file_type='XER',
            schedule_classification='CURRENT_UPDATE', data_date=date(2026, 9, 1),
            activities_json=acts, upload_timestamp=timezone.now() + timedelta(days=1),
        )
        section = report_service.build_float_intelligence_section(project, curr)
        self.assertEqual(len(section['topDeterioration']), 39)


class ProgressMilestoneSectionTests(TestCase):
    def setUp(self):
        self.project = Project.objects.create(name='Progress Milestone Report Test')
        now = timezone.now()
        self.baseline = ScheduleUpload.objects.create(
            project=self.project, original_filename='bl.xer', sanitized_filename='bl.xer', file_type='XER',
            schedule_classification='APPROVED_BASELINE', data_date=date(2026, 1, 1), version_label='Baseline',
            activities_json=[
                _act('A1', pctComplete=100.0),
                _act('M1', isMilestone=True, dur=0, bFinish='2026-01-15'),
            ],
            upload_timestamp=now - timedelta(days=30),
        )
        self.curr = ScheduleUpload.objects.create(
            project=self.project, original_filename='c.xer', sanitized_filename='c.xer', file_type='XER',
            schedule_classification='CURRENT_UPDATE', data_date=date(2026, 8, 15), version_label='Current',
            activities_json=[
                _act('A1', pctComplete=100.0),
                _act('M1', isMilestone=True, dur=0, bFinish='2026-01-25'),
            ],
            upload_timestamp=now,
        )

    def test_available_with_scurve_and_milestones(self):
        section = report_service.build_progress_milestone_section(self.project, self.curr)
        self.assertTrue(section['available'])
        self.assertIn('periods', section['scurve'])
        self.assertEqual(section['milestoneCount'], 1)
        self.assertEqual(section['slippedMilestoneCount'], 1)

    def test_unavailable_without_baseline(self):
        project2 = Project.objects.create(name='No Baseline Project')
        curr2 = ScheduleUpload.objects.create(
            project=project2, original_filename='c.xer', sanitized_filename='c.xer', file_type='XER',
            schedule_classification='CURRENT_UPDATE', data_date=date(2026, 8, 15),
            activities_json=[_act('A1')],
        )
        section = report_service.build_progress_milestone_section(project2, curr2)
        self.assertFalse(section['available'])


class FloatProgressExportTests(TestCase):
    def setUp(self):
        self.project = Project.objects.create(name='Float Progress Export Test')
        now = timezone.now()
        self.baseline = ScheduleUpload.objects.create(
            project=self.project, original_filename='bl.xer', sanitized_filename='bl.xer', file_type='XER',
            schedule_classification='APPROVED_BASELINE', data_date=date(2026, 1, 1), version_label='Baseline',
            activities_json=[
                _act('A1', totalFloat=20.0),
                _act('M1', isMilestone=True, dur=0, bFinish='2026-01-15', totalFloat=10.0),
            ],
            upload_timestamp=now - timedelta(days=30),
        )
        self.curr = ScheduleUpload.objects.create(
            project=self.project, original_filename='c.xer', sanitized_filename='c.xer', file_type='XER',
            schedule_classification='CURRENT_UPDATE', data_date=date(2026, 8, 15), version_label='Current',
            activities_json=[
                _act('A1', totalFloat=-5.0, isCritical=True),
                _act('M1', isMilestone=True, dur=0, bFinish='2026-01-25', totalFloat=-2.0),
            ],
            upload_timestamp=now,
        )
        self.payload = report_service.build_report_payload(self.project, 'WEEKLY_PROJECT_CONTROLS', version_id=str(self.curr.id))

    def test_pdf_renders_with_both_sections(self):
        pdf_bytes = report_export.generate_report_pdf(self.payload)
        self.assertTrue(pdf_bytes.startswith(b'%PDF'))

    def test_pdf_renders_without_sections(self):
        payload = dict(self.payload)
        payload['floatIntelligence'] = None
        payload['progressMilestones'] = None
        pdf_bytes = report_export.generate_report_pdf(payload)
        self.assertTrue(pdf_bytes.startswith(b'%PDF'))

    def test_excel_includes_all_new_worksheets(self):
        xlsx_bytes = report_export.generate_report_excel(self.payload)
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(xlsx_bytes))
        for name in (
            'Float Deterioration', 'Float Improvement', 'Newly Negative Float', 'Milestone Float',
            'Float Distribution', 'Progress Curve Data', 'Milestone Analysis',
        ):
            self.assertIn(name, wb.sheetnames, f'missing worksheet: {name}')

    def test_excel_float_deterioration_matches_real_row(self):
        xlsx_bytes = report_export.generate_report_excel(self.payload)
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(xlsx_bytes))
        rows = list(wb['Float Deterioration'].iter_rows(values_only=True))
        self.assertEqual(rows[1][0], 'A1')
        self.assertEqual(rows[1][5], -5.0)  # Current TF column
