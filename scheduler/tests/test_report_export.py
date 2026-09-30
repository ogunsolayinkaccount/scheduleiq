from django.test import SimpleTestCase

from scheduler.report_export import generate_report_excel, generate_report_pdf

_MINIMAL_PAYLOAD = {
    'engineVersion': '1.0.0', 'reportType': 'WEEKLY_PROJECT_CONTROLS',
    'sections': ['projectInfo', 'executiveSummary'],
    'projectInfo': {'projectName': 'Test Project', 'projectNumber': 'PN-1', 'dataDate': '2026-01-01', 'scheduleVersion': 'Jan Update'},
    'executiveSummary': {
        'narrative': {'sections': {
            'overallPerformance': 'Project cost performance is unfavorable with CPI 0.88 and current cost variance of -$1.2M.',
            'schedule': None, 'forecast': None, 'primaryDrivers': None, 'productivity': None,
        }},
        'aiNarrative': None,
        'controlSignals': [{'type': 'COMBINED_DETERIORATION', 'severity': 'critical', 'message': 'Both cost and schedule are unfavorable.'}],
        'warnings': [],
    },
    'currentPerformance': {
        'cpi': {'current': 0.88, 'previous': 0.92, 'delta': -0.04, 'pctChange': -4.3, 'direction': 'deteriorating'},
        'spi': {'current': 0.95, 'previous': 0.94, 'delta': 0.01, 'pctChange': 1.1, 'direction': 'improving'},
        'cv': {'current': -1_200_000.0, 'previous': -900_000.0, 'delta': None, 'pctChange': None, 'direction': 'deteriorating'},
        'sv': {'current': -100_000.0, 'previous': None, 'delta': None, 'pctChange': None, 'direction': 'unavailable'},
        'vac': {'current': -1_500_000.0, 'previous': None, 'delta': None, 'pctChange': None, 'direction': 'unavailable'},
    },
    'forecast': {
        'eacDrift': {
            'available': True,
            'scenarios': {
                'CPI_BASED': {'label': 'CPI-based', 'current': 11_800_000.0, 'previous': 11_500_000.0, 'delta': 300_000.0, 'direction': 'deteriorating'},
                'BOTTOM_UP': {'label': 'Bottom-up', 'current': 12_000_000.0, 'previous': 11_600_000.0, 'delta': 400_000.0, 'direction': 'deteriorating'},
            },
            'approvedEac': {'value': 12_400_000.0, 'description': 'PM approved Jan 2026'},
        },
        'forecastSignals': [],
    },
    'trendHistory': None,
    'schedulePerformance': {'available': True, 'classification': 'CURRENT_UPDATE', 'activityCount': 100, 'milestoneCount': 5, 'criticalCount': 20, 'negativeFloatCount': 2},
    'costProductivityDrivers': {
        'cost': {'drivers': [
            {'group': 'Area C', 'cv': -470_000.0, 'contributionPct': 39.2},
            {'group': 'Area D', 'cv': -310_000.0, 'contributionPct': 25.8},
        ]},
        'schedule': None, 'productivity': None,
    },
    'updateComparison': {
        'available': True, 'addedCount': 2, 'removedCount': 0, 'changedActivityCount': 15,
        'movedLaterCount': 5, 'movedEarlierCount': 1, 'newlyCriticalCount': 3, 'newlyNegativeFloatCount': 1,
        'floatDeterioratedCount': 4, 'floatImprovedCount': 1,
        'milestoneMovement': [{'activityName': 'Substantial Completion', 'previousFinish': '2026-06-01', 'currentFinish': '2026-06-12', 'deltaDays': 11}],
        'summaryNarrative': 'Substantial Completion slipped 11 calendar days this update.',
    },
    'managementAttention': [
        {'type': 'COST_VARIANCE', 'action': 'Investigate Area C — the largest unfavorable cost variance at -$470.0K.'},
    ],
    'dataQuality': ['No approved EAC has been entered for this project.'],
    'methodology': {'evMethod': 'DURATION_PCT_COMPLETE', 'pvMethod': 'LINEAR_BASELINE_SPREAD', 'cpiHealthyThreshold': 0.95, 'spiHealthyThreshold': 0.95, 'productivityHealthyThreshold': 0.90},
    'traceability': {'engineVersion': '1.0.0', 'evMethod': 'DURATION_PCT_COMPLETE', 'pvMethod': 'LINEAR_BASELINE_SPREAD', 'scheduleVersionId': 'abc', 'sourceUploadFilename': 'test.xer'},
}


class PdfExportTests(SimpleTestCase):
    def test_generates_valid_pdf_bytes(self):
        pdf_bytes = generate_report_pdf(_MINIMAL_PAYLOAD)
        self.assertTrue(pdf_bytes.startswith(b'%PDF'))
        self.assertGreater(len(pdf_bytes), 500)

    def test_handles_em_dash_and_unicode_punctuation(self):
        # Regression guard: fpdf2's core fonts raise on unicode em-dashes —
        # every narrative/action string in this app uses them routinely.
        payload = dict(_MINIMAL_PAYLOAD)
        payload['managementAttention'] = [{'type': 'X', 'action': 'Investigate Area C — CPI x SPI composite deteriorated.'}]
        pdf_bytes = generate_report_pdf(payload)
        self.assertTrue(pdf_bytes.startswith(b'%PDF'))

    def test_handles_missing_optional_sections(self):
        minimal = {
            'reportType': 'WEEKLY_PROJECT_CONTROLS', 'sections': [],
            'projectInfo': {}, 'executiveSummary': {}, 'currentPerformance': None,
            'forecast': {}, 'trendHistory': None, 'schedulePerformance': {'available': False},
            'costProductivityDrivers': None, 'updateComparison': None,
            'managementAttention': [], 'dataQuality': [], 'methodology': {}, 'traceability': {},
        }
        pdf_bytes = generate_report_pdf(minimal)
        self.assertTrue(pdf_bytes.startswith(b'%PDF'))

    def test_no_giant_activity_table(self):
        # The PDF must stay compact — no raw per-activity dump.
        pdf_bytes = generate_report_pdf(_MINIMAL_PAYLOAD)
        self.assertLess(len(pdf_bytes), 50_000)


class ExcelExportTests(SimpleTestCase):
    def test_generates_valid_xlsx_bytes(self):
        xlsx_bytes = generate_report_excel(_MINIMAL_PAYLOAD)
        self.assertTrue(xlsx_bytes.startswith(b'PK'))   # xlsx is a zip container
        self.assertGreater(len(xlsx_bytes), 1000)

    def test_handles_missing_optional_sections(self):
        minimal = {
            'reportType': 'WEEKLY_PROJECT_CONTROLS', 'sections': [],
            'projectInfo': {}, 'executiveSummary': {}, 'currentPerformance': None,
            'forecast': {}, 'trendHistory': None, 'schedulePerformance': {'available': False},
            'costProductivityDrivers': None, 'updateComparison': None,
            'managementAttention': [], 'dataQuality': [], 'methodology': {}, 'traceability': {},
        }
        xlsx_bytes = generate_report_excel(minimal)
        self.assertTrue(xlsx_bytes.startswith(b'PK'))

    def test_worksheets_present(self):
        from io import BytesIO
        from openpyxl import load_workbook
        xlsx_bytes = generate_report_excel(_MINIMAL_PAYLOAD)
        wb = load_workbook(BytesIO(xlsx_bytes))
        for expected in ('Executive Summary', 'Current Metrics', 'Trends', 'Forecast', 'Drivers', 'Schedule Changes', 'Milestones', 'Risks', 'Methodology'):
            self.assertIn(expected, wb.sheetnames)

    def test_unavailable_values_not_written_as_zero(self):
        from io import BytesIO
        from openpyxl import load_workbook
        xlsx_bytes = generate_report_excel(_MINIMAL_PAYLOAD)
        wb = load_workbook(BytesIO(xlsx_bytes))
        ws = wb['Current Metrics']
        sv_row = next(row for row in ws.iter_rows(min_row=2, values_only=True) if row[0] == 'SV')
        self.assertEqual(sv_row[2], 'Unavailable')   # 'previous' column, was None in payload


def _movement_row(activity_id, flags=('SLIPPED',)):
    return {
        'activityId': activity_id, 'activityName': f'Activity {activity_id}', 'wbs': 'Area A',
        'matchStatus': 'MATCHED', 'flags': list(flags),
        'previousStart': '2026-08-01', 'currentStart': '2026-08-01', 'startMovementDays': 0,
        'previousFinish': '2026-08-10', 'currentFinish': '2026-08-17', 'finishMovementDays': 7,
        'finishMovementWorkingDays': None, 'workingDayCalendarAvailable': False,
        'previousRemainingDuration': 10.0, 'currentRemainingDuration': 10.0, 'durationMovement': 0.0,
        'previousPctComplete': 0.0, 'currentPctComplete': 0.0,
        'previousTotalFloat': 0.0, 'currentTotalFloat': 0.0, 'floatMovement': 0.0,
        'previousStatus': 'NOT_STARTED', 'currentStatus': 'NOT_STARTED',
    }


_UPDATE_INTELLIGENCE_PAYLOAD = {
    'available': True,
    'previousVersionLabel': 'Update 02', 'previousDataDate': '2026-08-14',
    'currentVersionLabel': 'Update 03', 'currentDataDate': '2026-08-21', 'periodDays': 7,
    'movementRows': [_movement_row('A1')],
    'movementCounts': {'slipped': 1, 'improved': 0, 'unchanged': 0, 'new': 0, 'removed': 0, 'startedThisPeriod': 0, 'completedThisPeriod': 0},
    'reliability': {
        'available': True, 'plannedStarts': 2, 'actualStarts': 1, 'startReliabilityPct': 50.0,
        'plannedFinishes': 1, 'actualFinishes': 0, 'finishReliabilityPct': 0.0,
        'missedStarts': [{'activityId': 'A2', 'activityName': 'Missed Start', 'wbs': 'Area A', 'previousForecastStart': '2026-08-16', 'actualStart': None, 'status': 'FORECAST_START_MISSED'}],
        'missedFinishes': [{'activityId': 'A3', 'activityName': 'Missed Finish', 'wbs': 'Area A', 'previousForecastFinish': '2026-08-18', 'actualFinish': None, 'status': 'FORECAST_FINISH_MISSED'}],
    },
    'floatMovement': {'newlyCriticalCount': 1, 'newlyNegativeFloatCount': 0, 'leftCriticalPathCount': 0, 'recoveredFromNegativeFloatCount': 0},
    'logicChanges': {'added': [{'predecessorId': 'A0', 'successorId': 'A1', 'relType': 'FS', 'lagDays': 0}], 'removed': [], 'changed': []},
    'constraintChanges': [{'activityId': 'A1', 'activityName': 'Activity A1', 'wbs': 'Area A', 'constraintSlot': 'Primary', 'changeType': 'ADDED', 'previousType': None, 'currentType': 'CS_MEO', 'previousDate': None, 'currentDate': '2026-09-15', 'description': 'Primary constraint added: CS_MEO (2026-09-15)'}],
    'durationChanges': {'originalDurationChanges': [{'activityId': 'A1', 'activityName': 'Activity A1', 'wbs': 'Area A', 'previousDuration': 20.0, 'currentDuration': 30.0, 'delta': 10.0}], 'remainingDurationChanges': []},
    'milestoneMovement': {
        'rows': [{'activityId': 'MS1', 'activityName': 'Turnover', 'wbs': 'Area A', 'classification': 'SLIPPED', 'baselineDate': None, 'previousDate': '2026-09-01', 'currentDate': '2026-09-08', 'movementDays': 7, 'movementWorkingDays': None, 'workingDayCalendarAvailable': False, 'previousTotalFloat': 0.0, 'currentTotalFloat': 0.0}],
        'topSlippedMilestones': [{'activityName': 'Turnover', 'previousDate': '2026-09-01', 'currentDate': '2026-09-08', 'movementDays': 7, 'movementWorkingDays': None, 'workingDayCalendarAvailable': False}],
        'upcomingMilestones': [], 'milestonesCompletedThisPeriod': [], 'milestonesAtRisk': [],
    },
    'criticalPathMovement': {
        'methodologyNote': 'Critical/Driving Path Trace — not a ScheduleIQ CPM recalculation.',
        'stayedCritical': [], 'stayedCriticalCount': 0, 'becameCritical': [{'activityId': 'A1', 'activityName': 'Activity A1', 'wbs': 'Area A', 'currentStart': '2026-08-01', 'finishMovementDays': 7}], 'becameCriticalCount': 1,
        'leftCriticalPath': [], 'leftCriticalPathCount': 0, 'criticalActivitySlipped': [], 'criticalActivitySlippedCount': 0,
        'criticalFloatDeterioration': [], 'criticalFloatDeteriorationCount': 0,
        'criticalPathForecastDelayDays': 7, 'pathDivergence': {'activityId': 'A1', 'activityName': 'Activity A1', 'currentStart': '2026-08-01', 'note': 'x'},
    },
    'criticalPathShiftSummary': ['1 activities became newly critical.'],
    'driverAnalysis': {'groupBy': 'discipline', 'drivers': [{'group': 'Electrical', 'slippedCount': 1, 'cumulativeFinishMovementDays': 7, 'averageFinishMovementDays': 7.0, 'newlyNegativeFloatCount': 0, 'newlyCriticalCount': 1, 'missedForecastStartsCount': 0, 'missedForecastFinishesCount': 0, 'totalActivityCount': 5}], 'methodologyNote': 'aggregate movement indicator, not project completion delay'},
    'narrative': {'overview': '1 activity slipped, 0 improved, 0 started and 0 completed this update.', 'executionReliability': '2 starts were forecast.', 'criticalPath': '1 activity became newly critical.', 'milestones': '1 milestone slipped this update.', 'primaryDriver': 'Discipline Electrical contains the largest concentration of finish deterioration with 1 slipped activities.'},
    'hasBaseline': False,
    'lookaheadChange': {'available': True, 'previousWindow': {}, 'currentWindow': {}, 'carryover': [], 'carryoverCount': 0, 'newlyEntering': [], 'newlyEnteringCount': 0, 'pushedOut': [], 'pushedOutCount': 0, 'newlyCriticalNearTerm': [], 'newlyCriticalNearTermCount': 0},
}


class UpdateIntelligenceExportTests(SimpleTestCase):
    """Schedule Update Performance section — Weekly Look-Ahead Report
    Integration phase's own report_export.py additions."""

    def _payload(self, ui):
        return {**_MINIMAL_PAYLOAD, 'updateIntelligence': ui}

    def test_pdf_renders_update_performance_section(self):
        pdf_bytes = generate_report_pdf(self._payload(_UPDATE_INTELLIGENCE_PAYLOAD))
        self.assertTrue(pdf_bytes.startswith(b'%PDF'))

    def test_pdf_handles_unavailable_update_intelligence(self):
        pdf_bytes = generate_report_pdf(self._payload({'available': False, 'reason': 'No previous schedule version exists yet for this project.'}))
        self.assertTrue(pdf_bytes.startswith(b'%PDF'))

    def test_pdf_handles_missing_update_intelligence_key(self):
        payload = dict(_MINIMAL_PAYLOAD)
        pdf_bytes = generate_report_pdf(payload)   # no 'updateIntelligence' key at all
        self.assertTrue(pdf_bytes.startswith(b'%PDF'))

    def test_excel_worksheets_created(self):
        from io import BytesIO
        from openpyxl import load_workbook
        xlsx_bytes = generate_report_excel(self._payload(_UPDATE_INTELLIGENCE_PAYLOAD))
        wb = load_workbook(BytesIO(xlsx_bytes))
        for expected in (
            'Update Summary', 'Activity Movement', 'Missed Starts', 'Missed Finishes',
            'Critical Path Changes', 'Logic Changes', 'Duration Changes', 'Constraint Changes',
            'Update Milestones', 'Update Drivers', 'Look Ahead Changes',
        ):
            self.assertIn(expected, wb.sheetnames)

    def test_excel_no_duplicate_worksheet_names(self):
        from io import BytesIO
        from openpyxl import load_workbook
        xlsx_bytes = generate_report_excel(self._payload(_UPDATE_INTELLIGENCE_PAYLOAD))
        wb = load_workbook(BytesIO(xlsx_bytes))
        self.assertEqual(len(wb.sheetnames), len(set(wb.sheetnames)))
        # 'Update Milestones' must not collide with the pre-existing 'Milestones' sheet.
        self.assertIn('Milestones', wb.sheetnames)
        self.assertIn('Update Milestones', wb.sheetnames)

    def test_excel_skips_update_intelligence_sheets_when_unavailable(self):
        from io import BytesIO
        from openpyxl import load_workbook
        xlsx_bytes = generate_report_excel(self._payload({'available': False, 'reason': 'x'}))
        wb = load_workbook(BytesIO(xlsx_bytes))
        self.assertNotIn('Activity Movement', wb.sheetnames)
        self.assertIn('Milestones', wb.sheetnames)   # pre-existing sheets untouched

    def test_excel_activity_movement_full_detail_not_capped(self):
        """No-truncation regression: the Excel export must include EVERY
        applicable movement row, never a 2,500/5,000-style cap."""
        from io import BytesIO
        from openpyxl import load_workbook
        many_rows = [_movement_row(f'ACT-{i:05d}') for i in range(6000)]
        ui = dict(_UPDATE_INTELLIGENCE_PAYLOAD, movementRows=many_rows)
        xlsx_bytes = generate_report_excel(self._payload(ui))
        wb = load_workbook(BytesIO(xlsx_bytes))
        ws = wb['Activity Movement']
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(data_rows), 6000)
        # First, middle, AND last expected activity are all present — proves
        # no silent head-only slicing and no truncation at any point in the range.
        self.assertEqual(data_rows[0][0], 'ACT-00000')
        self.assertEqual(data_rows[3000][0], 'ACT-03000')
        self.assertEqual(data_rows[-1][0], 'ACT-05999')

    def test_excel_never_writes_fake_zero_for_unavailable_working_days(self):
        from io import BytesIO
        from openpyxl import load_workbook
        xlsx_bytes = generate_report_excel(self._payload(_UPDATE_INTELLIGENCE_PAYLOAD))
        wb = load_workbook(BytesIO(xlsx_bytes))
        ws = wb['Activity Movement']
        header = [c.value for c in ws[1]]
        idx = header.index('Finish Movement (working d)')
        first_data_row = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
        self.assertEqual(first_data_row[idx], 'Unavailable')
