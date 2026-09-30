"""
Baseline & Progress — Activity Chart Excel Export.

Unlike report_export.py (which reads ONLY a persisted, frozen
ProjectControlsReport snapshot — see that module's docstring), this export
is ad-hoc and live: it hands back exactly the current, filtered
baseline_progress.py row set the user is looking at in the Baseline vs
Forecast Activity Chart at the moment they click Export. It performs no
calculation of its own — every value comes straight from a row dict
baseline_progress.py already produced. Unavailable values are written as
the literal string 'Unavailable', never a blank-looking or fabricated 0,
matching the convention report_export.py already established.
"""

from __future__ import annotations

from datetime import date
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill(start_color='F0ECE4', end_color='F0ECE4', fill_type='solid')
_HEADER_FONT = Font(bold=True, size=10, color='5A4A3E')

_COLUMNS = (
    ('activityId', 'Activity ID'), ('activityName', 'Activity Name'), ('wbs', 'WBS'),
    ('baselineStart', 'Baseline Start'), ('baselineFinish', 'Baseline Finish'),
    ('currentStart', 'Current/Forecast Start'), ('currentFinish', 'Current/Forecast Finish'),
    ('startVarianceDays', 'Start Variance (calendar days)'), ('finishVarianceDays', 'Finish Variance (calendar days)'),
    ('startVarianceWorkingDays', 'Start Variance (working days)'), ('finishVarianceWorkingDays', 'Finish Variance (working days)'),
    ('pctComplete', '% Complete'), ('totalFloat', 'Total Float'), ('status', 'Status'),
)


def generate_activity_chart_excel(
    rows: List[Dict[str, Any]], data_date: Optional[date], project_name: str,
    current_version_label: Optional[str], baseline_version_label: Optional[str],
    scope_label: str = 'All Activities',
) -> bytes:
    """Builds a single-sheet workbook of the exact activity rows passed in —
    the same rows the chart's table/timeline are currently rendering, after
    global filters and Look-Ahead windowing have already been applied by
    the caller. Working-day variance columns are included alongside
    calendar-day ones (rather than picking one) so the export never hides
    which unit backs a given figure — the UI's wd/cd distinction stays
    fully traceable in the spreadsheet too."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Activity Chart'[:31]

    ws.append([f'Project: {project_name}'])
    ws.append([f'Current Version: {current_version_label or "Unavailable"}'])
    ws.append([f'Baseline Version: {baseline_version_label or "Unavailable — no baseline designated"}'])
    ws.append([f'Effective Data Date: {data_date.isoformat() if data_date else "Unavailable"}'])
    ws.append([f'Scope: {scope_label}'])
    ws.append([f'Activity Count: {len(rows)}'])
    ws.append([])
    for r in ws[1:6]:
        r[0].font = Font(bold=True, size=10, color='5A4A3E')

    header_row_idx = ws.max_row + 1
    ws.append([label for _, label in _COLUMNS])
    for c in ws[header_row_idx]:
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL

    for r in rows:
        out_row = []
        for key, _ in _COLUMNS:
            v = r.get(key)
            if key in ('startVarianceWorkingDays', 'finishVarianceWorkingDays') and not r.get('workingDayCalendarAvailable'):
                v = None  # never present a working-day figure the source calendar can't actually support
            out_row.append('Unavailable' if v is None else v)
        ws.append(out_row)

    for i in range(1, len(_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 20
    ws.column_dimensions['B'].width = 34

    import io
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
