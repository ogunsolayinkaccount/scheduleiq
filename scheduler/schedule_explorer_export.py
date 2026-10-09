"""
Schedule Explorer — Excel Export.

Live, ad-hoc export (same discipline as baseline_progress_export.py): it
hands back exactly the filtered activity rows the Schedule Explorer is
currently displaying at the moment the user clicks Export — nothing here
calculates anything; every value comes straight from a row dict
activity_analysis.py already produced. longestPathStatus is written
verbatim (YES/NO/UNAVAILABLE/UNKNOWN_LEGACY), never collapsed to a plain
Yes/No, so the export can never imply confirmed Longest Path membership
the screen itself doesn't claim.
"""

from __future__ import annotations

from datetime import date
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill(start_color='F0ECE4', end_color='F0ECE4', fill_type='solid')
_HEADER_FONT = Font(bold=True, size=10, color='5A4A3E')

_COLUMNS = (
    ('activityId', 'Activity ID'), ('activityName', 'Activity Name'),
    ('wbsPath', 'WBS'), ('area', 'Area'), ('activityStatus', 'Status'),
    ('currentStart', 'Start'), ('currentFinish', 'Finish'),
    ('originalDuration', 'Original Duration'), ('remainingDuration', 'Remaining Duration'),
    ('currentTotalFloat', 'Total Float'),
    ('criticalActionable', 'Critical Path'), ('longestPathStatus', 'Longest Path'),
)


def generate_schedule_explorer_excel(
    rows: List[Dict[str, Any]], project_name: str, current_version_label: Optional[str],
    data_date: Optional[date],
) -> bytes:
    """Builds a single-sheet workbook of exactly the activity rows passed
    in — the same rows the Explorer's own filters (WBS/area/status/search/
    path classification) already narrowed down, after the caller applied
    them. Never recalculates a date, Total Float, or path classification."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Schedule Explorer'[:31]

    ws.append([f'Project: {project_name}'])
    ws.append([f'Schedule Version: {current_version_label or "Unavailable"}'])
    ws.append([f'Effective Data Date: {data_date.isoformat() if data_date else "Unavailable"}'])
    ws.append([f'Activity Count: {len(rows)}'])
    ws.append([])
    for r in ws[1:4]:
        r[0].font = Font(bold=True, size=10, color='5A4A3E')

    header_row_idx = ws.max_row + 1
    ws.append([label for _, label in _COLUMNS])
    for c in ws[header_row_idx]:
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL

    for r in rows:
        out_row = []
        for key, _ in _COLUMNS:
            v = r.get(key)
            if key == 'criticalActionable':
                v = 'Critical' if v else 'Not Critical'
            out_row.append('Unavailable' if v is None else v)
        ws.append(out_row)

    for i in range(1, len(_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 20
    ws.column_dimensions['B'].width = 34
    ws.column_dimensions['C'].width = 30

    import io
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
