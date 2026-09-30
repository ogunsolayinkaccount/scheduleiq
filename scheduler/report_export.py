"""
Report Export — PDF & Excel — ScheduleIQ Automated Project Controls Reporting Phase

Both exporters read ONLY a persisted `ProjectControlsReport.payload_json`
snapshot — neither one calls cost_engine.py/trend_engine.py/driver_engine.py
or touches ScheduleUpload.activities_json. This is deliberate: the PDF/Excel
a user downloads must always match what they previewed and what's stored,
even if the calculation engines change in a later release. If a number
isn't already in the snapshot, it does not appear in the export.
"""

from __future__ import annotations

import io
from typing import Any, Dict, Optional

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


_ASCII_REPLACEMENTS = {
    '—': '-', '–': '-',       # em dash, en dash
    '‘': "'", '’': "'",       # curly single quotes
    '“': '"', '”': '"',       # curly double quotes
    '×': 'x',                       # multiplication sign (CPI x SPI)
    '≠': '!=', '≥': '>=', '≤': '<=',
    '▲': '^', '▼': 'v',       # trend arrows
    '→': '->', '…': '...',
    '•': '-',                       # bullet
}


def _ascii(text) -> str:
    """fpdf2's core (non-Unicode) fonts only support Latin-1 — this app's
    narrative/label text uses plain Unicode punctuation (em dashes, CPI x SPI,
    curly quotes) elsewhere in the UI/API. Rather than bundle a Unicode font
    file, report text is normalized to the closest ASCII equivalent here —
    values themselves are untouched, only punctuation characters."""
    s = str(text)
    for uni, ascii_eq in _ASCII_REPLACEMENTS.items():
        s = s.replace(uni, ascii_eq)
    return s.encode('latin-1', errors='replace').decode('latin-1')


def _narrative_text(exec_summary: Dict[str, Any]) -> Dict[str, Optional[str]]:
    """Prefers the AI-enhanced text when present and actually rewritten,
    falling back to the deterministic sentence per-section — the same
    fallback rule enhance_narrative_with_ai() already guarantees numeric
    integrity for. Exports never show a number the deterministic engine
    didn't produce."""
    deterministic = (exec_summary.get('narrative') or {}).get('sections') or {}
    ai = exec_summary.get('aiNarrative')
    if not ai:
        return deterministic
    ai_sections = ai.get('sections') or {}
    return {
        key: (ai_sections.get(key, {}).get('enhanced') if ai_sections.get(key, {}).get('wasRewritten') else text)
        for key, text in deterministic.items()
    }


def _money(v: Optional[float]) -> str:
    if v is None:
        return 'Unavailable'
    abs_v, sign = abs(v), '-' if v < 0 else ''
    if abs_v >= 1e6:
        return f'{sign}${abs_v / 1e6:.2f}M'
    if abs_v >= 1e3:
        return f'{sign}${abs_v / 1e3:.1f}K'
    return f'{sign}${abs_v:.0f}'


def _idx(v: Optional[float]) -> str:
    return 'Unavailable' if v is None else f'{v:.3f}'


def _num_or_dash(v) -> str:
    if v is None:
        return '-'
    if isinstance(v, float):
        return f'{v:.1f}'
    return str(v)


def _variance_display(row: Dict[str, Any]) -> str:
    """Prefers working-day variance when a confidently-decoded P6 calendar
    was available for that activity, otherwise falls back to calendar-day
    variance with the unit made explicit — never silently mixes the two."""
    days = row.get('finishVarianceDays')
    if days is None:
        return 'Unavailable'
    if row.get('workingDayCalendarAvailable') and row.get('finishVarianceWorkingDays') is not None:
        return f'{row["finishVarianceWorkingDays"]:+d} working days'
    return f'{days:+d} calendar days'


def _milestone_flags(m: Dict[str, Any]) -> str:
    flags = []
    if m.get('isDelayed'):
        flags.append('DELAYED')
    if m.get('isNegativeFloat'):
        flags.append('NEG FLOAT')
    if m.get('isApproaching'):
        flags.append('APPROACHING')
    return ' / '.join(flags) if flags else '-'


def _movement_display(days: Optional[int], working_days: Optional[int], cal_available: bool) -> str:
    """Same wd/cd-explicit convention as _variance_display(), for
    update_intelligence.py's movement rows (finishMovementDays /
    finishMovementWorkingDays / workingDayCalendarAvailable field names)."""
    if days is None:
        return 'Unavailable'
    if cal_available and working_days is not None:
        return f'{working_days:+d} working days'
    return f'{days:+d} calendar days'


def _pct_or_unavailable(v: Optional[float]) -> str:
    return 'Unavailable' if v is None else f'{v:.1f}%'


# ── PDF export ──────────────────────────────────────────────────────────────

class _ReportPDF(FPDF):
    def header(self):
        if getattr(self, '_suppress_header', False):
            return
        self.set_font('Helvetica', 'B', 9)
        self.set_text_color(120, 100, 84)
        self.cell(0, 6, 'ScheduleIQ - Project Controls Report', align='L')
        self.ln(10)

    def footer(self):
        self.set_y(-15)
        self.set_font('Helvetica', '', 8)
        self.set_text_color(150, 150, 150)
        self.cell(0, 10, f'Page {self.page_no()}', align='C')


def _section_title(pdf: FPDF, text: str):
    pdf.set_font('Helvetica', 'B', 13)
    pdf.set_text_color(30, 20, 16)
    pdf.cell(0, 9, _ascii(text), new_x='LMARGIN', new_y='NEXT')
    pdf.set_draw_color(210, 200, 190)
    pdf.line(pdf.get_x(), pdf.get_y(), pdf.get_x() + 190, pdf.get_y())
    pdf.ln(4)


def _kpi_row(pdf: FPDF, pairs):
    pdf.set_font('Helvetica', '', 10)
    col_w = 190 / len(pairs)
    y0 = pdf.get_y()
    for label, value in pairs:
        x0 = pdf.get_x()
        pdf.set_font('Helvetica', '', 8)
        pdf.set_text_color(120, 100, 84)
        pdf.multi_cell(col_w, 5, _ascii(label).upper(), align='L')
        pdf.set_xy(x0, pdf.get_y())
        pdf.set_font('Helvetica', 'B', 11)
        pdf.set_text_color(30, 20, 16)
        pdf.multi_cell(col_w, 6, _ascii(value), align='L')
        pdf.set_xy(x0 + col_w, y0)
    pdf.set_y(y0 + 16)


def _body_text(pdf: FPDF, text: str):
    pdf.set_font('Helvetica', '', 10)
    pdf.set_text_color(50, 40, 34)
    pdf.multi_cell(0, 5.5, _ascii(text))
    pdf.ln(1)


def _table(pdf: FPDF, headers, rows, col_widths=None):
    n = len(headers)
    col_widths = col_widths or [190 / n] * n
    pdf.set_font('Helvetica', 'B', 8)
    pdf.set_fill_color(240, 236, 228)
    pdf.set_text_color(90, 74, 62)
    for h, w in zip(headers, col_widths):
        pdf.cell(w, 6, _ascii(h), border=0, fill=True)
    pdf.ln()
    pdf.set_font('Helvetica', '', 8)
    pdf.set_text_color(40, 32, 26)
    for row in rows:
        for val, w in zip(row, col_widths):
            pdf.cell(w, 6, _ascii(val), border='B')
        pdf.ln()
    pdf.ln(3)


def generate_report_pdf(payload: Dict[str, Any]) -> bytes:
    """Builds the PDF entirely from `payload` (a persisted report snapshot
    dict) — see module docstring for why no calculation happens here."""
    pdf = _ReportPDF()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()

    info = payload.get('projectInfo') or {}
    report_type_label = 'Weekly Project Controls Report' if payload.get('reportType') == 'WEEKLY_PROJECT_CONTROLS' else 'Monthly Executive Report'

    pdf.set_font('Helvetica', 'B', 20)
    pdf.set_text_color(30, 20, 16)
    pdf.cell(0, 12, _ascii(info.get('projectName') or 'Project'), new_x='LMARGIN', new_y='NEXT')
    pdf.set_font('Helvetica', '', 12)
    pdf.set_text_color(90, 74, 62)
    pdf.cell(0, 8, report_type_label, new_x='LMARGIN', new_y='NEXT')
    pdf.set_font('Helvetica', '', 10)
    pdf.cell(0, 6, _ascii(f'Data Date: {info.get("dataDate") or "Unavailable"}   |   Schedule Version: {info.get("scheduleVersion") or "Unavailable"}'), new_x='LMARGIN', new_y='NEXT')
    pdf.ln(6)

    # ── Executive Controls Summary ──
    _section_title(pdf, 'Executive Controls Summary')
    exec_summary = payload.get('executiveSummary') or {}
    narrative = _narrative_text(exec_summary)
    overall = narrative.get('overallPerformance')
    _body_text(pdf, overall or 'Overall performance narrative unavailable — insufficient cost-loaded data.')

    perf = payload.get('currentPerformance') or {}
    _kpi_row(pdf, [
        ('CPI', _idx((perf.get('cpi') or {}).get('current'))),
        ('SPI', _idx((perf.get('spi') or {}).get('current'))),
        ('CV', _money((perf.get('cv') or {}).get('current'))),
        ('SV', _money((perf.get('sv') or {}).get('current'))),
    ])

    forecast = payload.get('forecast') or {}
    eac_drift = forecast.get('eacDrift') or {}
    scenarios = eac_drift.get('scenarios') or {}
    approved = eac_drift.get('approvedEac')
    _kpi_row(pdf, [
        ('Approved EAC', _money(approved['value']) if approved else 'None entered'),
        ('Bottom-Up EAC', _money((scenarios.get('BOTTOM_UP') or {}).get('current'))),
        ('VAC', _money((perf.get('vac') or {}).get('current'))),
        ('Productivity Factor', _idx((payload.get('trendHistory') or {}).get('productivity', {}).get('current') if payload.get('trendHistory') else None)),
    ])

    control_signals = exec_summary.get('controlSignals') or []
    if control_signals:
        pdf.set_font('Helvetica', 'B', 10)
        pdf.set_text_color(30, 20, 16)
        pdf.cell(0, 6, 'Control Signals', new_x='LMARGIN', new_y='NEXT')
        for s in control_signals:
            _body_text(pdf, f'- {s["message"]}')
    pdf.ln(2)

    # ── 4-Week Look Ahead ──
    lookahead = payload.get('lookAhead')
    if lookahead:
        _section_title(pdf, lookahead.get('scopeLabel') or '4-Week Look Ahead')
        if not lookahead.get('available'):
            _body_text(pdf, lookahead.get('reason') or 'Look-ahead data unavailable for this report.')
        else:
            window = lookahead.get('window') or {}
            _body_text(pdf, (
                f'Data Date: {lookahead.get("dataDate") or "Unavailable"}   |   '
                f'Look Ahead: {window.get("fromDate") or "Unavailable"} to {window.get("toDate") or "Unavailable"}   |   '
                f'Baseline: {lookahead.get("baselineVersionLabel") or "Unavailable"}   |   '
                f'Current: {lookahead.get("currentVersionLabel") or "Unavailable"}'
            ))

            cards = lookahead.get('summaryCards') or {}
            _kpi_row(pdf, [
                ('Activities in Window', cards.get('activitiesInWindow', 0)),
                ('Behind Baseline', cards.get('behindBaseline', 0)),
                ('Should Have Started', cards.get('shouldHaveStarted', 0)),
                ('Should Have Finished', cards.get('shouldHaveFinished', 0)),
            ])
            _kpi_row(pdf, [
                ('Planned Starts', cards.get('plannedStarts', 0)),
                ('Forecast Starts', cards.get('forecastStarts', 0)),
                ('In Progress', cards.get('inProgress', 0)),
                ('Critical Activities', cards.get('criticalActivities', 0)),
            ])
            if cards.get('plannedHours') is not None or cards.get('forecastHours') is not None or cards.get('actualHours') is not None:
                _kpi_row(pdf, [
                    ('Planned Hours', _num_or_dash(cards.get('plannedHours'))),
                    ('Forecast Hours', _num_or_dash(cards.get('forecastHours'))),
                    ('Actual Hours', _num_or_dash(cards.get('actualHours'))),
                    ('Avg Finish Variance', f'{cards["averageFinishVarianceDays"]:+.1f} d' if cards.get('averageFinishVarianceDays') is not None else 'Unavailable'),
                ])

            histogram = lookahead.get('histogram') or {}
            pdf.set_font('Helvetica', 'B', 10)
            pdf.set_text_color(30, 20, 16)
            pdf.cell(0, 6, f'Weekly Histogram ({(histogram.get("metric") or "activities").title()})', new_x='LMARGIN', new_y='NEXT')
            if histogram.get('available') and histogram.get('buckets'):
                _table(pdf, ['Period', 'Baseline', 'Current/Forecast', 'Actual'], [
                    (b['period'], _num_or_dash(b.get('baseline')), _num_or_dash(b.get('currentForecast')), _num_or_dash(b.get('actual')))
                    for b in histogram['buckets'][:8]
                ], col_widths=[46, 48, 58, 38])
            else:
                _body_text(pdf, histogram.get('reason') or 'No look-ahead window available.')

            scurve = lookahead.get('scurve') or {}
            pdf.set_font('Helvetica', 'B', 10)
            pdf.set_text_color(30, 20, 16)
            pdf.cell(0, 6, f'S-Curve ({(scurve.get("metric") or "duration").title()})', new_x='LMARGIN', new_y='NEXT')
            if scurve.get('available') and scurve.get('periods'):
                _table(pdf, ['Period', 'Baseline', 'Forecast', 'Actual', 'Status'], [
                    (p.get('period', ''), _num_or_dash(p.get('baselinePlanned')), _num_or_dash(p.get('currentPlanned')),
                     _num_or_dash(p.get('currentUpdateActual')), 'Past' if p.get('isPast') else 'Forecast')
                    for p in scurve['periods'][-8:]
                ], col_widths=[46, 36, 36, 36, 36])
            else:
                _body_text(pdf, scurve.get('reason') or 'S-Curve unavailable.')

            delayed = lookahead.get('delayedActivities') or []
            top_n = lookahead.get('topDelayedN', 10)
            pdf.set_font('Helvetica', 'B', 10)
            pdf.set_text_color(30, 20, 16)
            pdf.cell(0, 6, 'Top Delayed Activities', new_x='LMARGIN', new_y='NEXT')
            if delayed:
                _table(pdf, ['ID', 'Name', 'WBS', 'Baseline Fin.', 'Current Fin.', 'Variance', 'Status'], [
                    (d['activityId'], (d.get('activityName') or '')[:20], (d.get('wbs') or '')[:12],
                     d.get('baselineFinish') or '-', d.get('currentFinish') or '-', _variance_display(d), d.get('status', ''))
                    for d in delayed[:top_n]
                ], col_widths=[16, 44, 26, 26, 26, 30, 22])
                if len(delayed) > top_n:
                    _body_text(pdf, f'Showing top {top_n} of {len(delayed)} delayed activities.')
            else:
                _body_text(pdf, 'No delayed activities identified in the look-ahead window.')

            shs = lookahead.get('shouldHaveStarted') or []
            pdf.set_font('Helvetica', 'B', 10)
            pdf.set_text_color(30, 20, 16)
            pdf.cell(0, 6, 'Should Have Started', new_x='LMARGIN', new_y='NEXT')
            if shs:
                _table(pdf, ['ID', 'Name', 'Baseline Start', 'Days Overdue'], [
                    (d['activityId'], (d.get('activityName') or '')[:34], d.get('baselineStart') or '-', d.get('daysOverdue', '-'))
                    for d in shs
                ], col_widths=[20, 90, 40, 40])
                total = lookahead.get('shouldHaveStartedTotalCount', len(shs))
                if total > len(shs):
                    _body_text(pdf, f'Showing top {len(shs)} of {total}.')
            else:
                _body_text(pdf, 'No activities are should-have-started as of the Data Date.')

            shf = lookahead.get('shouldHaveFinished') or []
            pdf.set_font('Helvetica', 'B', 10)
            pdf.set_text_color(30, 20, 16)
            pdf.cell(0, 6, 'Should Have Finished', new_x='LMARGIN', new_y='NEXT')
            if shf:
                _table(pdf, ['ID', 'Name', 'Baseline Finish', 'Days Overdue'], [
                    (d['activityId'], (d.get('activityName') or '')[:34], d.get('baselineFinish') or '-', d.get('daysOverdue', '-'))
                    for d in shf
                ], col_widths=[20, 90, 40, 40])
                total = lookahead.get('shouldHaveFinishedTotalCount', len(shf))
                if total > len(shf):
                    _body_text(pdf, f'Showing top {len(shf)} of {total}.')
            else:
                _body_text(pdf, 'No activities are should-have-finished as of the Data Date.')

            milestones = lookahead.get('milestones') or []
            pdf.set_font('Helvetica', 'B', 10)
            pdf.set_text_color(30, 20, 16)
            pdf.cell(0, 6, 'Milestones in Window', new_x='LMARGIN', new_y='NEXT')
            if milestones:
                _table(pdf, ['Milestone', 'Baseline', 'Current', 'Variance', 'Float', 'Flags'], [
                    ((m.get('activityName') or '')[:28], m.get('baselineFinish') or '-', m.get('currentFinish') or '-',
                     _variance_display(m), _num_or_dash(m.get('totalFloat')), _milestone_flags(m))
                    for m in milestones
                ], col_widths=[46, 30, 30, 34, 20, 30])
            else:
                _body_text(pdf, 'No milestones fall within the look-ahead window.')
    pdf.ln(2)

    # ── Schedule Update Performance ──
    # Concise management summary only — the engine itself already analyzed
    # the COMPLETE schedule (see update_intelligence.py); this section picks
    # Top-N items for PDF readability, it does not re-limit the analysis.
    ui = payload.get('updateIntelligence')
    if ui:
        _section_title(pdf, 'Schedule Update Performance')
        if not ui.get('available'):
            _body_text(pdf, ui.get('reason') or 'Schedule Update Performance data unavailable for this report.')
        else:
            period_txt = f'{ui["periodDays"]} calendar days' if ui.get('periodDays') is not None else 'Unavailable'
            _body_text(pdf, (
                f'Previous: {ui.get("previousVersionLabel") or "Unavailable"} ({ui.get("previousDataDate") or "Unavailable"})   |   '
                f'Current: {ui.get("currentVersionLabel") or "Unavailable"} ({ui.get("currentDataDate") or "Unavailable"})   |   '
                f'Update Period: {period_txt}'
            ))

            reliability = ui.get('reliability') or {}
            _kpi_row(pdf, [
                ('Start Reliability', _pct_or_unavailable(reliability.get('startReliabilityPct'))),
                ('Finish Reliability', _pct_or_unavailable(reliability.get('finishReliabilityPct'))),
                ('Planned Starts', _num_or_dash(reliability.get('plannedStarts'))),
                ('Planned Finishes', _num_or_dash(reliability.get('plannedFinishes'))),
            ])

            mc = ui.get('movementCounts') or {}
            _kpi_row(pdf, [
                ('Slipped', mc.get('slipped', 0)), ('Improved', mc.get('improved', 0)),
                ('Started', mc.get('startedThisPeriod', 0)), ('Completed', mc.get('completedThisPeriod', 0)),
            ])

            float_result = ui.get('floatMovement') or {}
            cp_result = ui.get('criticalPathMovement') or {}
            _kpi_row(pdf, [
                ('Newly Critical', float_result.get('newlyCriticalCount', 0)),
                ('Newly Negative Float', float_result.get('newlyNegativeFloatCount', 0)),
                ('Left Critical Path', cp_result.get('leftCriticalPathCount', 0)),
                ('New / Removed', f'{mc.get("new", 0)} / {mc.get("removed", 0)}'),
            ])

            narrative = ui.get('narrative') or {}
            for key in ('overview', 'executionReliability', 'criticalPath', 'milestones', 'primaryDriver'):
                if narrative.get(key):
                    _body_text(pdf, narrative[key])

            missed_starts = reliability.get('missedStarts') or []
            if missed_starts:
                pdf.set_font('Helvetica', 'B', 10)
                pdf.set_text_color(30, 20, 16)
                pdf.cell(0, 6, 'Missed Forecast Starts', new_x='LMARGIN', new_y='NEXT')
                _table(pdf, ['ID', 'Name', 'WBS', 'Prev. Forecast Start'], [
                    (m['activityId'], (m.get('activityName') or '')[:34], (m.get('wbs') or '')[:20], m.get('previousForecastStart') or '-')
                    for m in missed_starts[:8]
                ], col_widths=[20, 76, 44, 50])
                if len(missed_starts) > 8:
                    _body_text(pdf, f'Showing top 8 of {len(missed_starts)} missed forecast starts.')

            missed_finishes = reliability.get('missedFinishes') or []
            if missed_finishes:
                pdf.set_font('Helvetica', 'B', 10)
                pdf.set_text_color(30, 20, 16)
                pdf.cell(0, 6, 'Missed Forecast Finishes', new_x='LMARGIN', new_y='NEXT')
                _table(pdf, ['ID', 'Name', 'WBS', 'Prev. Forecast Finish'], [
                    (m['activityId'], (m.get('activityName') or '')[:34], (m.get('wbs') or '')[:20], m.get('previousForecastFinish') or '-')
                    for m in missed_finishes[:8]
                ], col_widths=[20, 76, 44, 50])
                if len(missed_finishes) > 8:
                    _body_text(pdf, f'Showing top 8 of {len(missed_finishes)} missed forecast finishes.')

            drivers = (ui.get('driverAnalysis') or {}).get('drivers') or []
            top_drivers = [d for d in drivers if d.get('slippedCount', 0) > 0][:5]
            if top_drivers:
                pdf.set_font('Helvetica', 'B', 10)
                pdf.set_text_color(30, 20, 16)
                pdf.cell(0, 6, 'Top Deterioration Drivers', new_x='LMARGIN', new_y='NEXT')
                _table(pdf, ['Group', 'Slipped', 'Cumulative Movement', 'Newly Critical'], [
                    (d['group'], d['slippedCount'], f'{d["cumulativeFinishMovementDays"]:+d} activity-days', d.get('newlyCriticalCount', 0))
                    for d in top_drivers
                ], col_widths=[70, 40, 50, 30])

            top_milestones = (ui.get('milestoneMovement') or {}).get('topSlippedMilestones') or []
            if top_milestones:
                pdf.set_font('Helvetica', 'B', 10)
                pdf.set_text_color(30, 20, 16)
                pdf.cell(0, 6, 'Milestone Movement This Update', new_x='LMARGIN', new_y='NEXT')
                _table(pdf, ['Milestone', 'Previous Date', 'Current Date', 'Movement'], [
                    ((m.get('activityName') or '')[:38], m.get('previousDate') or '-', m.get('currentDate') or '-',
                     _movement_display(m.get('movementDays'), m.get('movementWorkingDays'), bool(m.get('workingDayCalendarAvailable'))))
                    for m in top_milestones[:5]
                ], col_widths=[60, 40, 40, 50])
    pdf.ln(2)

    # ── Float Intelligence ──
    fi = payload.get('floatIntelligence')
    if fi:
        _section_title(pdf, 'Float Intelligence')
        if not fi.get('available'):
            _body_text(pdf, fi.get('reason') or 'Float Intelligence data unavailable for this report.')
        else:
            s = fi.get('summary') or {}
            _kpi_row(pdf, [
                ('Negative Float', s.get('negativeFloatCount', 0)), ('Newly Negative', s.get('newlyNegativeCount', 0)),
                ('Deteriorated', s.get('floatDeterioratedCount', 0)), ('Improved', s.get('floatImprovedCount', 0)),
            ])
            _kpi_row(pdf, [
                ('Critical', s.get('criticalCount', 0)), ('Near Critical', s.get('nearCriticalCount', 0)),
                ('Median Current TF', _num_or_dash(s.get('medianCurrentTotalFloat'))),
                ('Minimum Current TF', _num_or_dash(s.get('minimumCurrentTotalFloat'))),
            ])
            top_det = (fi.get('topDeterioration') or [])[:8]
            if top_det:
                pdf.set_font('Helvetica', 'B', 10)
                pdf.set_text_color(30, 20, 16)
                pdf.cell(0, 6, 'Top Float Deterioration', new_x='LMARGIN', new_y='NEXT')
                _table(pdf, ['ID', 'Name', 'Baseline TF', 'Current TF', 'Delta', 'Finish'], [
                    (d['activityId'], (d.get('activityName') or '')[:26], _num_or_dash(d.get('baselineTotalFloat')),
                     _num_or_dash(d.get('currentTotalFloat')), _num_or_dash(d.get('floatChangeVsBaseline')), d.get('currentFinish') or '-')
                    for d in top_det
                ], col_widths=[20, 54, 30, 30, 26, 30])
                total_det = len(fi.get('topDeterioration') or [])
                if total_det > len(top_det):
                    _body_text(pdf, f'Showing top {len(top_det)} of {total_det} deteriorated activities — see Excel export for the complete detail.')
            milestone_exposure = [m for m in (fi.get('milestoneFloatAnalysis') or []) if (m.get('currentTotalFloat') or 0) < 0 or (m.get('finishVarianceDays') or 0) > 0]
            if milestone_exposure:
                pdf.set_font('Helvetica', 'B', 10)
                pdf.set_text_color(30, 20, 16)
                pdf.cell(0, 6, 'Milestone Float Exposure', new_x='LMARGIN', new_y='NEXT')
                _table(pdf, ['Milestone', 'Current Finish', 'Var (d)', 'Current TF'], [
                    ((m.get('activityName') or '')[:38], m.get('currentFinish') or '-', _num_or_dash(m.get('finishVarianceDays')), _num_or_dash(m.get('currentTotalFloat')))
                    for m in milestone_exposure[:8]
                ], col_widths=[70, 40, 30, 30])
    pdf.ln(2)

    # ── Progress & Milestones ──
    pm = payload.get('progressMilestones')
    if pm:
        _section_title(pdf, 'Progress & Milestones')
        if not pm.get('available'):
            _body_text(pdf, pm.get('reason') or 'Progress & Milestones data unavailable for this report.')
        else:
            scurve = pm.get('scurve') or {}
            _body_text(pdf, f'Data Date: {pm.get("currentDataDate") or "Unavailable"}  |  Weighting: {scurve.get("metric", "duration")}')
            periods = scurve.get('periods') or []
            past_periods = [p for p in periods if p.get('isPast')]
            latest_past = past_periods[-1] if past_periods else None
            _kpi_row(pdf, [
                ('Baseline Planned %', f'{latest_past["baselinePlanned"]}%' if latest_past and latest_past.get('baselinePlanned') is not None else 'Unavailable'),
                ('Actual/Earned %', f'{latest_past["currentUpdateActual"]}%' if latest_past and latest_past.get('currentUpdateActual') is not None else 'Unavailable'),
                ('Progress Variance', f'{latest_past["variance"]:+}%' if latest_past and latest_past.get('variance') is not None else 'Unavailable'),
                ('Forecast Completion', f'{periods[-1]["currentPlanned"]}%' if periods and periods[-1].get('currentPlanned') is not None else 'Unavailable'),
            ])
            _kpi_row(pdf, [
                ('Milestones', pm.get('milestoneCount', 0)), ('Upcoming', pm.get('upcomingMilestoneCount', 0)),
                ('Slipped', pm.get('slippedMilestoneCount', 0)), ('', ''),
            ])
            slipped = [m for m in (pm.get('milestones') or []) if (m.get('varianceDays') or 0) > 0][:8]
            if slipped:
                pdf.set_font('Helvetica', 'B', 10)
                pdf.set_text_color(30, 20, 16)
                pdf.cell(0, 6, 'Key Milestone Movements', new_x='LMARGIN', new_y='NEXT')
                _table(pdf, ['Milestone', 'Baseline', 'Current', 'Var (d)', 'Risk'], [
                    ((m.get('activityName') or '')[:34], m.get('baselineFinish') or '-', m.get('currentFinish') or '-',
                     _num_or_dash(m.get('varianceDays')), m.get('riskLevel') or '-')
                    for m in slipped
                ], col_widths=[64, 34, 34, 24, 24])
    pdf.ln(2)

    # ── Schedule Risk & Recovery ──
    # Concise management summary only — risk_register.py already analyzed the
    # COMPLETE risk population and recovery_engine.py the complete scenario;
    # this section shows Top-N for PDF readability without re-limiting the
    # underlying analysis (the Excel export carries the full detail).
    srr = payload.get('scheduleRiskRecovery')
    if srr:
        _section_title(pdf, 'Schedule Risk & Recovery')
        if not srr.get('available'):
            _body_text(pdf, srr.get('reason') or 'Schedule Risk & Recovery data unavailable for this report.')
        else:
            _body_text(pdf, f'Data Date: {srr.get("currentDataDate") or "Unavailable"}')
            summary = srr.get('summary') or {}
            by_sev = summary.get('bySeverity') or {}
            _kpi_row(pdf, [
                ('Critical Risks', by_sev.get('CRITICAL', 0)), ('High Risks', by_sev.get('HIGH', 0)),
                ('Driving', summary.get('drivingRiskCount', 0)), ('Milestone-Exposed', summary.get('milestoneExposureCount', 0)),
            ])
            accepted = srr.get('acceptedScenarios') or []
            _kpi_row(pdf, [
                ('Total Risks', summary.get('totalRisks', 0)), ('Accepted Scenarios', len(accepted)),
                ('Open Mitigation Actions', srr.get('openMitigationActionCount', 0)),
                ('Overdue Actions', srr.get('overdueMitigationActionCount', 0)),
            ])

            top_risks = [r for r in (srr.get('risks') or []) if r.get('severity') in ('CRITICAL', 'HIGH')][:8]
            if top_risks:
                pdf.set_font('Helvetica', 'B', 10)
                pdf.set_text_color(30, 20, 16)
                pdf.cell(0, 6, 'Top Schedule Risks', new_x='LMARGIN', new_y='NEXT')
                _table(pdf, ['ID', 'Name', 'Severity', 'Urgency', 'Driving', 'Reason'], [
                    (r['activityId'], (r.get('activityName') or '')[:24], r['severity'], r['urgency'],
                     'Y' if r.get('driving') else '-', (r.get('severityReason') or '')[:40])
                    for r in top_risks
                ], col_widths=[20, 50, 24, 24, 20, 62])
                total_risks = summary.get('totalRisks', len(srr.get('risks') or []))
                if total_risks > len(top_risks):
                    _body_text(pdf, f'Showing top {len(top_risks)} Critical/High risks of {total_risks} total — see Excel export for the complete register.')
            else:
                _body_text(pdf, 'No Critical or High severity risks identified for this update.')

            if accepted:
                pdf.set_font('Helvetica', 'B', 10)
                pdf.set_text_color(30, 20, 16)
                pdf.cell(0, 6, 'Accepted Recovery Scenarios', new_x='LMARGIN', new_y='NEXT')
                _table(pdf, ['Scenario', 'Risk', 'Recovery (days)', 'Baseline Variance Recovered'], [
                    (s['name'][:34], s.get('riskKey') or '-', _num_or_dash(s.get('recoveryDays')),
                     _num_or_dash((s.get('projectBaselineRecovery') or {}).get('varianceRecoveredDays')))
                    for s in accepted[:8]
                ], col_widths=[70, 30, 40, 60])

            overdue = [a for a in (srr.get('mitigationActions') or []) if a.get('overdue')][:8]
            if overdue:
                pdf.set_font('Helvetica', 'B', 10)
                pdf.set_text_color(180, 40, 40)
                pdf.cell(0, 6, 'Overdue Mitigation Actions', new_x='LMARGIN', new_y='NEXT')
                pdf.set_text_color(30, 20, 16)
                _table(pdf, ['Description', 'Risk', 'Owner', 'Due Date'], [
                    (a['description'][:44], a.get('riskKey') or '-', a.get('owner') or '-', a.get('dueDate') or '-')
                    for a in overdue
                ], col_widths=[80, 30, 40, 30])

            overdue_count = srr.get('overdueMitigationActionCount', 0)
            overdue_clause = f', {overdue_count} overdue' if overdue_count else ''
            narrative_line = (
                f'{summary.get("totalRisks", 0)} risk(s) identified this update, '
                f'{by_sev.get("CRITICAL", 0)} Critical and {by_sev.get("HIGH", 0)} High. '
                f'{len(accepted)} recovery scenario(s) accepted as the working mitigation plan '
                f'(this does not change the imported P6 schedule). '
                f'{srr.get("openMitigationActionCount", 0)} mitigation action(s) open{overdue_clause}.'
            )
            _body_text(pdf, narrative_line)
    pdf.ln(2)

    # ── Forecast ──
    _section_title(pdf, 'Forecast')
    if scenarios:
        _table(pdf, ['Methodology', 'Current EAC', 'Previous EAC', 'Direction'], [
            (s.get('label', k), _money(s.get('current')), _money(s.get('previous')), s.get('direction', 'unavailable'))
            for k, s in scenarios.items()
        ], col_widths=[70, 45, 45, 30])
    else:
        _body_text(pdf, 'No EAC forecast scenarios available — project is not cost-loaded.')
    if approved:
        _body_text(pdf, f'Approved EAC: {_money(approved["value"])} ({approved.get("description") or "no description"})')

    # ── Key Drivers ──
    all_drivers = payload.get('costProductivityDrivers') or {}
    if all_drivers.get('cost') and all_drivers['cost'].get('drivers'):
        _section_title(pdf, 'Key Cost Drivers')
        _table(pdf, ['Group', 'CV', 'Contribution'], [
            (d['group'], _money(d['cv']), f'{d["contributionPct"]}%' if d.get('contributionPct') is not None else '—')
            for d in all_drivers['cost']['drivers'][:8]
        ], col_widths=[90, 55, 45])

    # ── Schedule Update ──
    update_comp = payload.get('updateComparison')
    if update_comp and update_comp.get('available'):
        _section_title(pdf, 'Schedule Update')
        _body_text(pdf, update_comp.get('summaryNarrative') or '')
        _kpi_row(pdf, [
            ('Added', update_comp['addedCount']), ('Removed', update_comp['removedCount']),
            ('Moved Later', update_comp['movedLaterCount']), ('Newly Critical', update_comp['newlyCriticalCount']),
        ])

    sched_perf = payload.get('schedulePerformance') or {}
    if sched_perf.get('available'):
        _kpi_row(pdf, [
            ('Critical Activities', sched_perf.get('criticalCount')),
            ('Negative Float Count', sched_perf.get('negativeFloatCount')),
            ('Milestones', sched_perf.get('milestoneCount')),
            ('Classification', sched_perf.get('classification') or 'Unavailable'),
        ])

    # ── Management Attention ──
    actions = payload.get('managementAttention') or []
    _section_title(pdf, 'Management Attention')
    if actions:
        for a in actions:
            _body_text(pdf, f'- {a["action"]}')
    else:
        _body_text(pdf, 'No management-attention items identified from currently available data.')

    # ── Methodology / Data Notes ──
    _section_title(pdf, 'Methodology & Data Notes')
    methodology = payload.get('methodology') or {}
    trace = payload.get('traceability') or {}
    _body_text(pdf, (
        f'EV Method: {methodology.get("evMethod", "n/a")}   |   PV Method: {methodology.get("pvMethod", "n/a")}   |   '
        f'Engine Version: {trace.get("engineVersion", "n/a")}'
    ))
    data_quality = payload.get('dataQuality') or []
    for item in data_quality:
        _body_text(pdf, f'- {item}')

    out = pdf.output()
    return bytes(out)


# ── Excel export ─────────────────────────────────────────────────────────────

_HEADER_FILL = PatternFill(start_color='F0ECE4', end_color='F0ECE4', fill_type='solid')
_HEADER_FONT = Font(bold=True, size=10, color='5A4A3E')


def _write_kv_sheet(ws, title: str, rows):
    ws.title = title[:31]
    ws.append(['Field', 'Value'])
    for c in ws[1]:
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
    for label, value in rows:
        ws.append([label, value if value is not None else 'Unavailable'])
    ws.column_dimensions['A'].width = 32
    ws.column_dimensions['B'].width = 50


def _write_table_sheet(ws, title: str, headers, rows):
    ws.title = title[:31]
    ws.append(headers)
    for c in ws[1]:
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
    for row in rows:
        ws.append(['Unavailable' if v is None else v for v in row])
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 20


def generate_report_excel(payload: Dict[str, Any]) -> bytes:
    """Builds the workbook entirely from `payload` — same no-recompute
    rule as generate_report_pdf(). Unavailable values are written as the
    literal string 'Unavailable', never a blank-looking 0."""
    wb = Workbook()

    info = payload.get('projectInfo') or {}
    exec_summary = payload.get('executiveSummary') or {}
    narrative_sections = _narrative_text(exec_summary)

    ws = wb.active
    _write_kv_sheet(ws, 'Executive Summary', [
        ('Project', info.get('projectName')), ('Project Number', info.get('projectNumber')),
        ('Data Date', info.get('dataDate')), ('Previous Update', info.get('previousUpdateDate')),
        ('Report Type', payload.get('reportType')),
        ('Overall Performance', narrative_sections.get('overallPerformance')),
        ('Schedule', narrative_sections.get('schedule')),
        ('Forecast', narrative_sections.get('forecast')),
        ('Primary Drivers', narrative_sections.get('primaryDrivers')),
        ('Productivity', narrative_sections.get('productivity')),
    ])

    perf = payload.get('currentPerformance') or {}
    ws2 = wb.create_sheet('Current Metrics')
    _write_table_sheet(ws2, 'Current Metrics', ['Metric', 'Current', 'Previous', 'Direction'], [
        (k.upper(), v.get('current'), v.get('previous'), v.get('direction')) for k, v in perf.items()
    ] if perf else [])

    trends = payload.get('trendHistory')
    ws3 = wb.create_sheet('Trends')
    if trends:
        _write_table_sheet(ws3, 'Trends', ['Metric', 'Direction', 'Consecutive Count', 'Observation'], [
            (name, t.get('direction'), t.get('consecutiveCount'), t.get('observation'))
            for name, t in trends.items()
        ])
    else:
        _write_kv_sheet(ws3, 'Trends', [('Status', 'Unavailable — fewer than 2 persisted schedule versions')])

    forecast = payload.get('forecast') or {}
    scenarios = (forecast.get('eacDrift') or {}).get('scenarios') or {}
    approved = (forecast.get('eacDrift') or {}).get('approvedEac')
    ws4 = wb.create_sheet('Forecast')
    _write_table_sheet(ws4, 'Forecast', ['Methodology', 'Current EAC', 'Previous EAC', 'Delta', 'Direction'], [
        (s.get('label', k), s.get('current'), s.get('previous'), s.get('delta'), s.get('direction'))
        for k, s in scenarios.items()
    ])
    if approved:
        ws4.append([])
        ws4.append(['Approved EAC', approved['value'], '', '', approved.get('description') or ''])

    drivers = payload.get('costProductivityDrivers') or {}
    ws5 = wb.create_sheet('Drivers')
    row_cursor = 1
    for dim in ('cost', 'schedule', 'productivity'):
        d = drivers.get(dim)
        if not d or not d.get('drivers'):
            continue
        ws5.cell(row=row_cursor, column=1, value=f'{dim.upper()} DRIVERS').font = Font(bold=True)
        row_cursor += 1
        metric_key = {'cost': 'cv', 'schedule': 'sv', 'productivity': 'productivityFactor'}[dim]
        ws5.append(['Group', metric_key.upper(), 'Contribution %'])
        row_cursor += 1
        for item in d['drivers']:
            ws5.append([item['group'], item.get(metric_key), item.get('contributionPct')])
            row_cursor += 1
        row_cursor += 1
    ws5.column_dimensions['A'].width = 24

    # ── 4-Week Look Ahead worksheets (all filtered rows — never truncated
    # to match the PDF's Top 10, per this section's own docstring rule) ──
    lookahead = payload.get('lookAhead')
    if lookahead and lookahead.get('available'):
        cards = lookahead.get('summaryCards') or {}
        window = lookahead.get('window') or {}
        ws_summary = wb.create_sheet('Look Ahead Summary')
        _write_kv_sheet(ws_summary, 'Look Ahead Summary', [
            ('Data Date', lookahead.get('dataDate')),
            ('Look Ahead Start', window.get('fromDate')), ('Look Ahead Finish', window.get('toDate')),
            ('Baseline Version', lookahead.get('baselineVersionLabel')), ('Current Version', lookahead.get('currentVersionLabel')),
            ('Filters', ', '.join(f'{k}={v}' for k, v in (lookahead.get('filters') or {}).items()) or 'All'),
            ('Activities in Window', cards.get('activitiesInWindow')),
            ('Planned Starts', cards.get('plannedStarts')), ('Forecast Starts', cards.get('forecastStarts')),
            ('Planned Finishes', cards.get('plannedFinishes')), ('Forecast Finishes', cards.get('forecastFinishes')),
            ('Behind Baseline', cards.get('behindBaseline')),
            ('Should Have Started', cards.get('shouldHaveStarted')), ('Should Have Finished', cards.get('shouldHaveFinished')),
            ('In Progress', cards.get('inProgress')), ('Critical Activities', cards.get('criticalActivities')),
            ('Average Finish Variance (days)', cards.get('averageFinishVarianceDays')),
            ('Planned Hours', cards.get('plannedHours')), ('Forecast Hours', cards.get('forecastHours')), ('Actual Hours', cards.get('actualHours')),
            ('Histogram Metric', (lookahead.get('histogram') or {}).get('metric')),
            ('S-Curve Metric', (lookahead.get('scurve') or {}).get('metric')),
        ])

        rows = lookahead.get('rows') or []
        ws_activities = wb.create_sheet('4-Week Look Ahead')
        _write_table_sheet(ws_activities, '4-Week Look Ahead', [
            'Activity ID', 'Name', 'WBS', 'Area', 'Discipline', 'Contractor', 'System',
            'Baseline Start', 'Baseline Finish', 'Current Start', 'Current Finish',
            'Actual Start', 'Actual Finish', '% Complete', 'Total Float',
            'Start Variance (d)', 'Finish Variance (d)', 'Finish Variance (working d)', 'Status',
        ], [
            (r.get('activityId'), r.get('activityName'), r.get('wbs'), r.get('area'), r.get('discipline'), r.get('contractor'), r.get('system'),
             r.get('baselineStart'), r.get('baselineFinish'), r.get('currentStart'), r.get('currentFinish'),
             r.get('actualStart'), r.get('actualFinish'), r.get('pctComplete'), r.get('totalFloat'),
             r.get('startVarianceDays'), r.get('finishVarianceDays'),
             r.get('finishVarianceWorkingDays') if r.get('workingDayCalendarAvailable') else None, r.get('status'))
            for r in rows
        ])

        delayed = lookahead.get('delayedActivities') or []
        ws_delayed = wb.create_sheet('Delayed Activities')
        _write_table_sheet(ws_delayed, 'Delayed Activities', [
            'Activity ID', 'Name', 'WBS', 'Area', 'Baseline Finish', 'Current Finish',
            'Finish Variance (d)', 'Finish Variance (working d)', 'Total Float', '% Complete', 'Status',
        ], [
            (d.get('activityId'), d.get('activityName'), d.get('wbs'), d.get('area'), d.get('baselineFinish'), d.get('currentFinish'),
             d.get('finishVarianceDays'), d.get('finishVarianceWorkingDays') if d.get('workingDayCalendarAvailable') else None,
             d.get('totalFloat'), d.get('pctComplete'), d.get('status'))
            for d in delayed
        ])

        ws_shs = wb.create_sheet('Should Have Started')
        _write_table_sheet(ws_shs, 'Should Have Started', ['Activity ID', 'Name', 'WBS', 'Baseline Start', 'Days Overdue'], [
            (d.get('activityId'), d.get('activityName'), d.get('wbs'), d.get('baselineStart'), d.get('daysOverdue'))
            for d in (lookahead.get('shouldHaveStarted') or [])
        ])
        ws_shf = wb.create_sheet('Should Have Finished')
        _write_table_sheet(ws_shf, 'Should Have Finished', ['Activity ID', 'Name', 'WBS', 'Baseline Finish', 'Days Overdue'], [
            (d.get('activityId'), d.get('activityName'), d.get('wbs'), d.get('baselineFinish'), d.get('daysOverdue'))
            for d in (lookahead.get('shouldHaveFinished') or [])
        ])

        ws_la_milestones = wb.create_sheet('LookAhead Milestones')
        _write_table_sheet(ws_la_milestones, 'LookAhead Milestones', [
            'Milestone', 'Baseline Date', 'Current Date', 'Finish Variance (d)', 'Total Float', 'Delayed', 'Negative Float', 'Approaching',
        ], [
            (m.get('activityName'), m.get('baselineFinish'), m.get('currentFinish'), m.get('finishVarianceDays'), m.get('totalFloat'),
             m.get('isDelayed'), m.get('isNegativeFloat'), m.get('isApproaching'))
            for m in (lookahead.get('milestones') or [])
        ])

    # ── Schedule Update Performance (update_intelligence.py) — FULL detail,
    # never truncated to match the PDF's Top-N management view. Every row
    # here comes straight from the persisted payload; nothing is
    # recomputed or resampled at export time. ──
    ui = payload.get('updateIntelligence')
    if ui and ui.get('available'):
        reliability = ui.get('reliability') or {}
        mc = ui.get('movementCounts') or {}
        cp = ui.get('criticalPathMovement') or {}
        fm = ui.get('floatMovement') or {}

        ws_uis = wb.create_sheet('Update Summary')
        _write_kv_sheet(ws_uis, 'Update Summary', [
            ('Previous Version', ui.get('previousVersionLabel')), ('Previous Data Date', ui.get('previousDataDate')),
            ('Current Version', ui.get('currentVersionLabel')), ('Current Data Date', ui.get('currentDataDate')),
            ('Update Period (calendar days)', ui.get('periodDays')),
            ('Start Reliability %', reliability.get('startReliabilityPct')),
            ('Planned Starts', reliability.get('plannedStarts')), ('Actual Starts', reliability.get('actualStarts')),
            ('Finish Reliability %', reliability.get('finishReliabilityPct')),
            ('Planned Finishes', reliability.get('plannedFinishes')), ('Actual Finishes', reliability.get('actualFinishes')),
            ('Slipped', mc.get('slipped')), ('Improved', mc.get('improved')), ('Unchanged', mc.get('unchanged')),
            ('New', mc.get('new')), ('Removed', mc.get('removed')),
            ('Started This Period', mc.get('startedThisPeriod')), ('Completed This Period', mc.get('completedThisPeriod')),
            ('Newly Critical', fm.get('newlyCriticalCount')), ('Left Critical Path', fm.get('leftCriticalPathCount')),
            ('Newly Negative Float', fm.get('newlyNegativeFloatCount')), ('Recovered From Negative Float', fm.get('recoveredFromNegativeFloatCount')),
            ('Critical Path Forecast Delay (days)', cp.get('criticalPathForecastDelayDays')),
            ('Path Divergence Activity', (cp.get('pathDivergence') or {}).get('activityId')),
            ('Overview', (ui.get('narrative') or {}).get('overview')),
            ('Execution Reliability', (ui.get('narrative') or {}).get('executionReliability')),
            ('Critical Path', (ui.get('narrative') or {}).get('criticalPath')),
            ('Milestones', (ui.get('narrative') or {}).get('milestones')),
            ('Primary Driver', (ui.get('narrative') or {}).get('primaryDriver')),
        ])

        movement_rows = ui.get('movementRows') or []
        ws_movement = wb.create_sheet('Activity Movement')
        _write_table_sheet(ws_movement, 'Activity Movement', [
            'Activity ID', 'Activity Name', 'WBS', 'Match Status', 'Flags',
            'Previous Start', 'Current Start', 'Start Movement (d)',
            'Previous Finish', 'Current Finish', 'Finish Movement (d)', 'Finish Movement (working d)',
            'Previous Remaining Duration', 'Current Remaining Duration', 'Duration Movement',
            'Previous % Complete', 'Current % Complete',
            'Previous Total Float', 'Current Total Float', 'Float Movement',
            'Previous Status', 'Current Status', 'Baseline Finish', 'Baseline Variance (d)',
        ], [
            (r.get('activityId'), r.get('activityName'), r.get('wbs'), r.get('matchStatus'), ', '.join(r.get('flags') or []),
             r.get('previousStart'), r.get('currentStart'), r.get('startMovementDays'),
             r.get('previousFinish'), r.get('currentFinish'), r.get('finishMovementDays'),
             r.get('finishMovementWorkingDays') if r.get('workingDayCalendarAvailable') else None,
             r.get('previousRemainingDuration'), r.get('currentRemainingDuration'), r.get('durationMovement'),
             r.get('previousPctComplete'), r.get('currentPctComplete'),
             r.get('previousTotalFloat'), r.get('currentTotalFloat'), r.get('floatMovement'),
             r.get('previousStatus'), r.get('currentStatus'), r.get('baselineFinish'), r.get('baselineVarianceDays'))
            for r in movement_rows
        ])

        commitment = ui.get('forecastCommitment') or {}
        ws_missed_starts = wb.create_sheet('Missed Starts')
        _write_table_sheet(ws_missed_starts, 'Missed Starts', ['Activity ID', 'Activity Name', 'WBS', 'Previous Forecast Start', 'Actual Start', 'Status'], [
            (s.get('activityId'), s.get('activityName'), s.get('wbs'), s.get('previousForecastStart'), s.get('actualStart'), s.get('status'))
            for s in (commitment.get('startCommitments') or [])
        ])
        ws_missed_finishes = wb.create_sheet('Missed Finishes')
        _write_table_sheet(ws_missed_finishes, 'Missed Finishes', ['Activity ID', 'Activity Name', 'WBS', 'Previous Forecast Finish', 'Actual Finish', 'Status'], [
            (f.get('activityId'), f.get('activityName'), f.get('wbs'), f.get('previousForecastFinish'), f.get('actualFinish'), f.get('status'))
            for f in (commitment.get('finishCommitments') or [])
        ])

        ws_cp = wb.create_sheet('Critical Path Changes')
        cp_row = 1
        for label, key in (('STAYED CRITICAL', 'stayedCritical'), ('BECAME CRITICAL', 'becameCritical'),
                            ('LEFT CRITICAL PATH', 'leftCriticalPath'), ('CRITICAL ACTIVITY SLIPPED', 'criticalActivitySlipped'),
                            ('CRITICAL FLOAT DETERIORATION', 'criticalFloatDeterioration')):
            entries = cp.get(key) or []
            ws_cp.cell(row=cp_row, column=1, value=label).font = Font(bold=True)
            cp_row += 1
            ws_cp.append(['Activity ID', 'Activity Name', 'WBS', 'Current Start', 'Finish Movement (d)'])
            cp_row += 1
            for e in entries:
                ws_cp.append([e.get('activityId'), e.get('activityName'), e.get('wbs'), e.get('currentStart'), e.get('finishMovementDays')])
                cp_row += 1
            cp_row += 1
        ws_cp.column_dimensions['A'].width = 22
        ws_cp.column_dimensions['B'].width = 30

        logic = ui.get('logicChanges') or {}
        ws_logic = wb.create_sheet('Logic Changes')
        _write_table_sheet(ws_logic, 'Logic Changes', ['Change Type', 'Predecessor ID', 'Successor ID', 'Field', 'Previous', 'Current'], [
            (kind.upper(), e.get('predecessorId'), e.get('successorId'), e.get('field', ''), e.get('previous', e.get('relType')), e.get('current', e.get('lagDays')))
            for kind in ('added', 'removed', 'changed') for e in (logic.get(kind) or [])
        ])

        ws_duration = wb.create_sheet('Duration Changes')
        duration = ui.get('durationChanges') or {}
        dur_rows = (
            [('ORIGINAL', d.get('activityId'), d.get('activityName'), d.get('wbs'), d.get('previousDuration'), d.get('currentDuration'), d.get('delta'))
             for d in (duration.get('originalDurationChanges') or [])]
            + [('REMAINING', d.get('activityId'), d.get('activityName'), d.get('wbs'), d.get('previousDuration'), d.get('currentDuration'), d.get('delta'))
               for d in (duration.get('remainingDurationChanges') or [])]
        )
        _write_table_sheet(ws_duration, 'Duration Changes', ['Type', 'Activity ID', 'Activity Name', 'WBS', 'Previous Duration', 'Current Duration', 'Delta'], dur_rows)

        ws_constraint = wb.create_sheet('Constraint Changes')
        _write_table_sheet(ws_constraint, 'Constraint Changes', ['Activity ID', 'Activity Name', 'WBS', 'Slot', 'Change Type', 'Previous Type', 'Current Type', 'Previous Date', 'Current Date', 'Description'], [
            (c.get('activityId'), c.get('activityName'), c.get('wbs'), c.get('constraintSlot'), c.get('changeType'),
             c.get('previousType'), c.get('currentType'), c.get('previousDate'), c.get('currentDate'), c.get('description'))
            for c in (ui.get('constraintChanges') or [])
        ])

        ms = ui.get('milestoneMovement') or {}
        ws_ms = wb.create_sheet('Update Milestones')
        _write_table_sheet(ws_ms, 'Update Milestones', [
            'Activity ID', 'Activity Name', 'WBS', 'Classification', 'Baseline Date', 'Previous Date', 'Current Date',
            'Baseline Variance (d)', 'This Update Movement (d)', 'This Update Movement (working d)',
            'Previous Total Float', 'Current Total Float',
        ], [
            (m.get('activityId'), m.get('activityName'), m.get('wbs'), m.get('classification'), m.get('baselineDate'),
             m.get('previousDate'), m.get('currentDate'), m.get('baselineVarianceDays'), m.get('movementDays'),
             m.get('movementWorkingDays') if m.get('workingDayCalendarAvailable') else None,
             m.get('previousTotalFloat'), m.get('currentTotalFloat'))
            for m in (ms.get('rows') or [])
        ])

        driver_result = ui.get('driverAnalysis') or {}
        ws_drivers = wb.create_sheet('Update Drivers')
        _write_table_sheet(ws_drivers, 'Update Drivers', [
            'Group', 'Slipped Count', 'Cumulative Finish Movement (activity-days)', 'Average Finish Movement (d)',
            'Newly Negative Float', 'Newly Critical', 'Missed Forecast Starts', 'Missed Forecast Finishes', 'Total Activities in Group',
        ], [
            (d.get('group'), d.get('slippedCount'), d.get('cumulativeFinishMovementDays'), d.get('averageFinishMovementDays'),
             d.get('newlyNegativeFloatCount'), d.get('newlyCriticalCount'), d.get('missedForecastStartsCount'),
             d.get('missedForecastFinishesCount'), d.get('totalActivityCount'))
            for d in (driver_result.get('drivers') or [])
        ])
        if driver_result.get('methodologyNote'):
            ws_drivers.append([])
            ws_drivers.append(['Methodology', driver_result['methodologyNote']])

        lac = ui.get('lookaheadChange') or {}
        ws_lac = wb.create_sheet('Look Ahead Changes')
        if lac.get('available'):
            lac_row = 1
            for label, key in (('CARRYOVER', 'carryover'), ('NEWLY ENTERING', 'newlyEntering'),
                                ('PUSHED OUT', 'pushedOut'), ('NEWLY CRITICAL NEAR-TERM', 'newlyCriticalNearTerm')):
                entries = lac.get(key) or []
                ws_lac.cell(row=lac_row, column=1, value=label).font = Font(bold=True)
                lac_row += 1
                ws_lac.append(['Activity ID', 'Activity Name', 'WBS', 'Baseline Finish', 'Current Finish', 'Status'])
                lac_row += 1
                for e in entries:
                    ws_lac.append([e.get('activityId'), e.get('activityName'), e.get('wbs'), e.get('baselineFinish'), e.get('currentFinish'), e.get('status')])
                    lac_row += 1
                lac_row += 1
            ws_lac.column_dimensions['A'].width = 22
            ws_lac.column_dimensions['B'].width = 30
        else:
            _write_kv_sheet(ws_lac, 'Look Ahead Changes', [('Status', lac.get('reason', 'Unavailable'))])

    # ── Float Intelligence (activity_analysis.py / float_intelligence.py) —
    # FULL detail, never Top-N truncated. ──
    fi = payload.get('floatIntelligence')
    if fi and fi.get('available'):
        ws_float_det = wb.create_sheet('Float Deterioration')
        _write_table_sheet(ws_float_det, 'Float Deterioration', [
            'Activity ID', 'Activity Name', 'WBS', 'Baseline TF', 'Previous TF', 'Current TF',
            'Delta vs Baseline', 'Delta vs Previous', 'Finish Variance (d)', 'Current Finish', 'Critical', 'Driving', 'Milestone',
        ], [
            (d.get('activityId'), d.get('activityName'), d.get('wbs'), d.get('baselineTotalFloat'), d.get('previousTotalFloat'),
             d.get('currentTotalFloat'), d.get('floatChangeVsBaseline'), d.get('floatChangeVsPrevious'),
             d.get('finishVarianceDays'), d.get('currentFinish'), d.get('critical'), d.get('driving'), d.get('isMilestone'))
            for d in (fi.get('topDeterioration') or [])
        ])

        ws_float_imp = wb.create_sheet('Float Improvement')
        _write_table_sheet(ws_float_imp, 'Float Improvement', [
            'Activity ID', 'Activity Name', 'WBS', 'Baseline TF', 'Previous TF', 'Current TF',
            'Delta vs Baseline', 'Delta vs Previous', 'Finish Variance (d)', 'Current Finish', 'Critical', 'Driving',
        ], [
            (d.get('activityId'), d.get('activityName'), d.get('wbs'), d.get('baselineTotalFloat'), d.get('previousTotalFloat'),
             d.get('currentTotalFloat'), d.get('floatChangeVsBaseline'), d.get('floatChangeVsPrevious'),
             d.get('finishVarianceDays'), d.get('currentFinish'), d.get('critical'), d.get('driving'))
            for d in (fi.get('topImprovement') or [])
        ])

        ws_newly_neg = wb.create_sheet('Newly Negative Float')
        _write_table_sheet(ws_newly_neg, 'Newly Negative Float', [
            'Activity ID', 'Activity Name', 'WBS', 'Previous TF', 'Current TF', 'Finish Variance (d)', 'Current Finish',
        ], [
            (d.get('activityId'), d.get('activityName'), d.get('wbs'), d.get('previousTotalFloat'),
             d.get('currentTotalFloat'), d.get('finishVarianceDays'), d.get('currentFinish'))
            for d in (fi.get('newlyNegative') or [])
        ])

        ws_milestone_float = wb.create_sheet('Milestone Float')
        _write_table_sheet(ws_milestone_float, 'Milestone Float', [
            'Milestone ID', 'Milestone Name', 'Baseline Finish', 'Current Finish', 'Finish Variance (d)',
            'Baseline TF', 'Previous TF', 'Current TF', 'Delta vs Baseline', 'Driving Predecessors', 'Negative-Float Predecessors', 'Risk Flagged',
        ], [
            (m.get('activityId'), m.get('activityName'), m.get('baselineFinish'), m.get('currentFinish'), m.get('finishVarianceDays'),
             m.get('baselineTotalFloat'), m.get('previousTotalFloat'), m.get('currentTotalFloat'), m.get('floatChangeVsBaseline'),
             m.get('drivingPredecessorCount'), m.get('negativeFloatPredecessorExposure'), m.get('riskFlagged'))
            for m in (fi.get('milestoneFloatAnalysis') or [])
        ])

        ws_float_hist = wb.create_sheet('Float Distribution')
        _write_table_sheet(ws_float_hist, 'Float Distribution', ['Range', 'Current', 'Baseline', 'Previous'], [
            (b.get('range'), b.get('current'), b.get('baseline'), b.get('previous'))
            for b in (fi.get('distribution') or {}).get('buckets', [])
        ])
    elif fi:
        ws_float_det = wb.create_sheet('Float Deterioration')
        _write_kv_sheet(ws_float_det, 'Float Deterioration', [('Status', fi.get('reason', 'Unavailable'))])

    # ── Progress & Milestones (progress_curve.py / baseline_progress.py /
    # milestones.py) — FULL detail, never Top-N truncated. ──
    pm = payload.get('progressMilestones')
    if pm and pm.get('available'):
        scurve = pm.get('scurve') or {}
        ws_curve = wb.create_sheet('Progress Curve Data')
        _write_table_sheet(ws_curve, 'Progress Curve Data', [
            'Period', 'Period End', 'Is Past (Data Date)', 'Baseline Planned %', 'Actual/Earned %',
            'Current Forecast %', 'Variance', 'Incremental Planned', 'Incremental Actual',
        ], [
            (p.get('period'), p.get('periodEnd'), p.get('isPast'), p.get('baselinePlanned'), p.get('currentUpdateActual'),
             p.get('currentPlanned'), p.get('variance'), p.get('incrementalPlanned'), p.get('incrementalActual'))
            for p in (scurve.get('periods') or [])
        ])
        ws_curve.append([])
        ws_curve.append(['Weighting Basis', scurve.get('metric')])
        ws_curve.append(['Methodology', scurve.get('methodologyNote')])

        ws_ms_analysis = wb.create_sheet('Milestone Analysis')
        _write_table_sheet(ws_ms_analysis, 'Milestone Analysis', [
            'Activity ID', 'Activity Name', 'WBS', 'Status', 'Baseline Finish', 'Current Finish', 'Variance (d)',
            'Total Float', 'Critical', 'Risk Level', 'Driving Predecessor', 'Previous Update Finish', 'Movement Since Last Update (d)',
        ], [
            (m.get('activityId'), m.get('activityName'), m.get('wbs'), m.get('status'), m.get('baselineFinish'), m.get('currentFinish'),
             m.get('varianceDays'), m.get('totalFloat'), m.get('isCritical'), m.get('riskLevel'),
             (m.get('drivingPredecessor') or {}).get('activityId'), m.get('previousUpdateFinish'), m.get('movementSinceLastUpdateDays'))
            for m in (pm.get('milestones') or [])
        ])
    elif pm:
        ws_curve = wb.create_sheet('Progress Curve Data')
        _write_kv_sheet(ws_curve, 'Progress Curve Data', [('Status', pm.get('reason', 'Unavailable'))])

    # ── Schedule Risk & Recovery (risk_register.py / recovery_engine.py /
    # driving_chain.py) — FULL detail, never Top-N truncated. Every risk,
    # scenario, scenario assumption, mitigation action, and recovery
    # tracking comparison in the persisted snapshot is written here. ──
    srr = payload.get('scheduleRiskRecovery')
    if srr and srr.get('available'):
        ws_risks = wb.create_sheet('Schedule Risks')
        _write_table_sheet(ws_risks, 'Schedule Risks', [
            'Activity ID', 'Activity Name', 'WBS', 'Area', 'Discipline', 'Severity', 'Severity Reason',
            'Urgency', 'Urgency Reason', 'Driving', 'Critical', 'Total Float', 'Baseline Finish',
            'Baseline Variance (d)', 'Current Finish', 'Update Movement (d)', 'Risk Signals',
            'Connected Milestones', 'Forecast-Impacting Milestones', 'Status', 'Owner', 'Target Date', 'Mitigation Notes',
        ], [
            (r.get('activityId'), r.get('activityName'), r.get('wbs'), r.get('area'), r.get('discipline'),
             r.get('severity'), r.get('severityReason'), r.get('urgency'), r.get('urgencyReason'),
             r.get('driving'), r.get('critical'), r.get('totalFloat'), r.get('baselineFinish'), r.get('baselineVarianceDays'),
             r.get('currentFinish'), r.get('updateMovementDays'),
             '; '.join(s.get('description', '') for s in (r.get('riskSignals') or [])),
             (r.get('milestoneExposure') or {}).get('connectedMilestoneCount'),
             len((r.get('milestoneExposure') or {}).get('forecastImpactingMilestones') or []),
             (r.get('workflow') or {}).get('status'), (r.get('workflow') or {}).get('owner'),
             (r.get('workflow') or {}).get('targetDate'), (r.get('workflow') or {}).get('mitigationNotes'))
            for r in (srr.get('risks') or [])
        ])

        ws_scenarios = wb.create_sheet('Recovery Scenarios')
        _write_table_sheet(ws_scenarios, 'Recovery Scenarios', [
            'Scenario', 'Status', 'Linked Risk', 'Source Version', 'Source Data Date', 'Action Count',
            'Current Forecast Finish', 'Scenario Forecast Finish', 'Recovery (days)',
            'Baseline Finish', 'Current Baseline Variance (d)', 'Scenario Baseline Variance (d)',
            'Variance Recovered (d)', 'Recovery %', 'Improved', 'Worsened', 'Newly Critical', 'Newly Negative Float',
        ], [
            (s.get('name'), s.get('status'), s.get('riskKey'), s.get('sourceVersionLabel'), s.get('sourceDataDate'), s.get('actionCount'),
             s.get('currentForecastFinish'), s.get('scenarioForecastFinish'), s.get('recoveryDays'),
             (s.get('projectBaselineRecovery') or {}).get('baselineFinish'),
             (s.get('projectBaselineRecovery') or {}).get('currentBaselineVarianceDays'),
             (s.get('projectBaselineRecovery') or {}).get('scenarioBaselineVarianceDays'),
             (s.get('projectBaselineRecovery') or {}).get('varianceRecoveredDays'),
             (s.get('projectBaselineRecovery') or {}).get('recoveryPct'),
             (s.get('sideEffects') or {}).get('improvedCount'), (s.get('sideEffects') or {}).get('worsenedCount'),
             (s.get('sideEffects') or {}).get('newlyCriticalCount'), (s.get('sideEffects') or {}).get('newlyNegativeFloatCount'))
            for s in (srr.get('scenarios') or [])
        ])

        ws_scenario_changes = wb.create_sheet('Scenario Changes')
        change_rows = []
        for s in (srr.get('scenarios') or []):
            for a in (s.get('assumptions') or []):
                change_rows.append((
                    s.get('name'), a.get('type'), a.get('activityId'), a.get('predecessorId'), a.get('successorId'),
                    a.get('newDuration'), a.get('newLagDays'), a.get('relType') or a.get('newRelType'), a.get('lagDays'), a.get('note'),
                ))
        _write_table_sheet(ws_scenario_changes, 'Scenario Changes', [
            'Scenario', 'Action Type', 'Activity ID', 'Predecessor ID', 'Successor ID',
            'New Duration', 'New Lag (d)', 'Relationship Type', 'Lag (d)', 'Note',
        ], change_rows)

        ws_mitigation = wb.create_sheet('Mitigation Actions')
        _write_table_sheet(ws_mitigation, 'Mitigation Actions', [
            'Description', 'Linked Risk', 'Linked Scenario', 'Owner', 'Due Date', 'Status', 'Overdue', 'Notes',
        ], [
            (a.get('description'), a.get('riskKey'), a.get('scenarioId'), a.get('owner'), a.get('dueDate'),
             a.get('status'), a.get('overdue'), a.get('notes'))
            for a in (srr.get('mitigationActions') or [])
        ])

        ws_tracking = wb.create_sheet('Recovery Tracking')
        tracking_rows = []
        for s in (srr.get('scenarios') or []):
            rt = s.get('recoveryTracking')
            if not rt or not rt.get('available'):
                continue
            for c in rt.get('milestoneComparisons') or []:
                tracking_rows.append((
                    s.get('name'), rt.get('nextVersionLabel'), rt.get('nextDataDate'),
                    c.get('activityId'), c.get('activityName'), c.get('classification'),
                    c.get('scenarioTargetFinish'), c.get('nextUpdateForecastFinish'), c.get('differenceDays'),
                ))
        _write_table_sheet(ws_tracking, 'Recovery Tracking', [
            'Scenario', 'Compared Against Version', 'Compared Against Data Date', 'Milestone ID', 'Milestone Name',
            'Classification', 'Scenario Target Finish', 'Next Update Forecast Finish', 'Difference (d)',
        ], tracking_rows)

        ws_milestone_recovery = wb.create_sheet('Milestone Recovery')
        milestone_recovery_rows = []
        for s in (srr.get('scenarios') or []):
            for m in (s.get('milestoneImpact') or []):
                milestone_recovery_rows.append((
                    s.get('name'), s.get('status'), m.get('activityId'), m.get('activityName'),
                    m.get('currentForecastFinish'), m.get('scenarioForecastFinish'), m.get('movementDays'),
                    m.get('baselineFinish'), m.get('currentBaselineVarianceDays'), m.get('scenarioBaselineVarianceDays'),
                    m.get('varianceRecoveredDays'), m.get('recoveryPct'),
                ))
        _write_table_sheet(ws_milestone_recovery, 'Milestone Recovery', [
            'Scenario', 'Scenario Status', 'Milestone ID', 'Milestone Name', 'Current Forecast Finish',
            'Scenario Forecast Finish', 'Movement (d)', 'Baseline Finish', 'Current Baseline Variance (d)',
            'Scenario Baseline Variance (d)', 'Variance Recovered (d)', 'Recovery %',
        ], milestone_recovery_rows)
    elif srr:
        ws_risks = wb.create_sheet('Schedule Risks')
        _write_kv_sheet(ws_risks, 'Schedule Risks', [('Status', srr.get('reason', 'Unavailable'))])

    update_comp = payload.get('updateComparison')
    ws6 = wb.create_sheet('Schedule Changes')
    if update_comp and update_comp.get('available'):
        _write_kv_sheet(ws6, 'Schedule Changes', [
            ('Added', update_comp['addedCount']), ('Removed', update_comp['removedCount']),
            ('Changed', update_comp['changedActivityCount']), ('Moved Later', update_comp['movedLaterCount']),
            ('Moved Earlier', update_comp['movedEarlierCount']), ('Newly Critical', update_comp['newlyCriticalCount']),
            ('Newly Negative Float', update_comp['newlyNegativeFloatCount']),
            ('Summary', update_comp.get('summaryNarrative')),
        ])
    else:
        _write_kv_sheet(ws6, 'Schedule Changes', [('Status', (update_comp or {}).get('reason', 'Unavailable'))])

    ws7 = wb.create_sheet('Milestones')
    milestones = (update_comp or {}).get('milestoneMovement') or []
    _write_table_sheet(ws7, 'Milestones', ['Activity', 'Previous Finish', 'Current Finish', 'Delta (days)'], [
        (m.get('activityName'), m.get('previousFinish'), m.get('currentFinish'), m.get('deltaDays')) for m in milestones
    ])

    ws8 = wb.create_sheet('Risks')
    signals = (exec_summary.get('controlSignals') or []) + (forecast.get('forecastSignals') or [])
    actions = payload.get('managementAttention') or []
    _write_table_sheet(ws8, 'Risks', ['Type', 'Message'], [(s.get('type'), s.get('message')) for s in signals])
    ws8.append([])
    ws8.append(['MANAGEMENT ATTENTION', ''])
    for a in actions:
        ws8.append([a.get('type'), a.get('action')])

    ws9 = wb.create_sheet('Methodology')
    methodology = payload.get('methodology') or {}
    trace = payload.get('traceability') or {}
    _write_kv_sheet(ws9, 'Methodology', [
        ('EV Method', methodology.get('evMethod')), ('PV Method', methodology.get('pvMethod')),
        ('CPI Healthy Threshold', methodology.get('cpiHealthyThreshold')),
        ('SPI Healthy Threshold', methodology.get('spiHealthyThreshold')),
        ('Productivity Healthy Threshold', methodology.get('productivityHealthyThreshold')),
        ('Engine Version', trace.get('engineVersion')), ('Source Upload', trace.get('sourceUploadFilename')),
        ('Schedule Version ID', trace.get('scheduleVersionId')),
        *[(f'Data Quality Note {i+1}', item) for i, item in enumerate(payload.get('dataQuality') or [])],
    ])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
